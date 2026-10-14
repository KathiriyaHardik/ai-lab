"""Fast tests for everything except the AI itself (a fake Ollama client stands in for it).

    uv run pytest

The real model is tested separately by evaluate.py.
"""
import io
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import ollama
import pytest

HERE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HERE))
import agent  # noqa: E402

JOB_AD = (HERE / "samples" / "job_ai_engineer_en.txt").read_text()
TRIES = 1 + agent.MAX_FIX_ROUNDS  # a first draft plus every fix round
CV = (HERE / "samples" / "cv_ai_student.txt").read_text()


class FakeClient:
    """Plays back prepared replies (strings or exceptions) and records every call."""

    def __init__(self, *replies):
        self.replies = list(replies)
        self.calls = []

    def chat(self, **kwargs):
        self.calls.append(kwargs)
        reply = self.replies.pop(0)
        if isinstance(reply, Exception):
            raise reply
        # the agent streams: send the reply in small pieces, like Ollama does
        return iter(SimpleNamespace(message=SimpleNamespace(content=reply[i:i + 7])) for i in range(0, len(reply), 7))


def strength(point, evidence, requirement="Strong Python skills"):
    return {"point": point, "requirement": requirement, "evidence": evidence}


def plan(*strengths, name="ALEX MUSTER", gap="", gap_answer=""):
    strengths = strengths or (
        strength("RAG", "Built a RAG chatbot in Python with LangChain"),
        strength("Docker", "Packaged the service with Docker"),
    )
    return json.dumps({
        "candidate_name": name, "company": "Brightline Analytics GmbH", "job_title": "Junior AI Engineer (m/f/d)",
        "contact_person": "", "strengths": list(strengths), "gap": gap, "gap_answer": gap_answer,
        "motivation": "Builds AI assistants for manufacturers",
    })


GOOD_PARAGRAPHS = [
    "Brightline Analytics builds AI assistants that help manufacturers find answers in their own documents, and "
    "that is exactly the kind of work I want to do as a Junior AI Engineer.",
    "As a working student I built a RAG chatbot in Python with LangChain and a Chroma vector database that answers "
    "questions from 2,000 internal PDF pages. I connected it to the OpenAI API and to local models in Ollama, and I "
    "packaged the service with Docker and deployed it on an internal server. I have used Python for 4 years.",
    "I have not worked with Kubernetes yet, but I learned Docker on the job and I pick up new tools quickly. My "
    "pytest tests and code reviews on GitHub show that I care about software that keeps working after the demo.",
    "In my projects I also built a prompt testing notebook that compares answers from three different LLMs. Measuring "
    "answer quality is part of your role description, and it is the part of AI engineering I enjoy most, because it "
    "turns a nice demo into a tool that people can actually trust.",
    "I would be glad to show you the chatbot and talk about how I could help your team in an interview.",
]


def draft(paragraphs=None, greeting="Dear Hiring Team,", closing="Kind regards\nAlex Muster",
          subject="Application: Junior AI Engineer at Brightline Analytics"):
    return json.dumps({"subject": subject, "greeting": greeting, "paragraphs": paragraphs or GOOD_PARAGRAPHS,
                       "closing": closing})


def cut_reply(*ids):
    """The AI's answer in the cut step: which sentences to remove."""
    return json.dumps({"remove": list(ids)})


def planned(kind="cover_letter", language="English", notes="", **plan_args):
    app = agent.start(JOB_AD, CV, kind, language, notes)
    return agent.make_plan(app, FakeClient(plan(**plan_args)))


def problems(app=None, **draft_args):
    app = app or planned()
    return agent.check_draft(app, agent.Draft.model_validate_json(draft(**draft_args)))


# ----------------------------------------------------------------------------- small helpers
def test_core_names():
    assert agent.core_name("Brightline Analytics GmbH") == "brightline analytics"
    assert agent.core_name("Junior AI Engineer (m/f/d)") == "junior ai engineer"
    assert agent.core_name("Müller & Co. KG") == "müller"


def test_numbers_ignore_thousand_separators():
    assert agent.numbers("2,000 pages, 2.000 Seiten, 4 years, 12/345") == {"2000", "4", "12", "345"}


# ----------------------------------------------------------------------------- Python's checks on a draft
def test_a_good_draft_passes():
    assert problems() == []


def test_too_short_and_too_long():
    assert "too short" in problems(paragraphs=["I would like to work at Brightline Analytics as a Junior AI Engineer."])[0]
    assert "too long" in problems(paragraphs=GOOD_PARAGRAPHS * 2)[0]


def test_message_has_its_own_length():
    app = planned(kind="message")
    short = ["I build RAG chatbots in Python and would love to do that as a Junior AI Engineer at Brightline "
             "Analytics. Could we have a short call next week?"] * 2
    assert problems(app, paragraphs=short, greeting="Hello,") == []
    assert "too long" in problems(app, greeting="Hello,")[0]


def test_company_and_job_title_must_be_mentioned():
    found = problems(paragraphs=[p.replace("Brightline Analytics", "your company").replace("Junior AI Engineer", "engineer")
                                 for p in GOOD_PARAGRAPHS], subject="Application")
    assert 'Mention the company "Brightline Analytics GmbH".' in found
    assert 'Mention the job title "Junior AI Engineer (m/f/d)".' in found


def test_greeting_closing_and_name():
    found = problems(greeting="Yo team,", closing="Cheers\nSomeone")
    assert any(p.startswith("Start with a greeting") for p in found)
    assert any(p.startswith("End with 'Kind regards'") for p in found)
    assert 'Sign with the name "Alex Muster".' in found


def test_german_greeting_and_closing():
    app = planned(language="Deutsch")
    found = problems(app, greeting="Dear Hiring Team,", closing="Kind regards\nAlex Muster")
    assert any("Sehr geehrte" in p for p in found) and any("Mit freundlichen Grüßen" in p for p in found)


def test_placeholders_are_caught():
    found = problems(paragraphs=GOOD_PARAGRAPHS[:-1] + ["I look forward to hearing from [Hiring Manager] soon in an interview."])
    assert 'Replace the placeholder "[Hiring Manager]" with real content.' in found


def test_invented_numbers_are_caught_but_real_ones_pass():
    invented = [p.replace("for 4 years", "for 7 years").replace("2,000", "5,000") for p in GOOD_PARAGRAPHS]
    found = problems(paragraphs=invented)
    assert any("not in the CV, the job ad or the notes: 7, 5000" in p for p in found)
    app = planned(notes="I can start on 1 March 2027.")
    assert problems(app, paragraphs=GOOD_PARAGRAPHS[:-1] + ["I can start on 1 March 2027 and would be glad to talk in an interview."]) == []


def test_cliches_and_wrong_language():
    found = problems(paragraphs=["I am writing to apply for the job."] + GOOD_PARAGRAPHS[1:])
    assert any("empty phrase" in p for p in found)
    german = ["Ich habe mit Python und LangChain einen RAG Chatbot für die Firma gebaut und ich bin sehr motiviert, "
              "die Arbeit als Junior AI Engineer bei Brightline Analytics zu machen."] * 4
    assert "Write the whole text in English." in problems(paragraphs=german)


# ----------------------------------------------------------------------------- starting and planning
@pytest.mark.parametrize("job_ad, cv, message", [("", CV, "job ad is empty"), (JOB_AD, " ", "CV is empty"),
                                                  (JOB_AD, CV * 30, "too long")])
def test_bad_inputs(job_ad, cv, message):
    with pytest.raises(agent.AgentError, match=message):
        agent.start(job_ad, cv)


def test_plan_proof_is_checked_against_the_cv():
    app = agent.make_plan(agent.start(JOB_AD, CV), FakeClient(plan(
        strength("RAG", "Built a RAG chatbot in Python with LangChain"),
        strength("Kubernetes", "Ran production clusters on Kubernetes"),  # not in the CV
        strength("rag ", "Built a RAG chatbot in Python with LangChain"),  # the same point again
        strength("Docker", "Packaged the service with Docker"),
        strength("Git", "Git and GitHub"),
        strength("pytest", "pytest"),
    )))
    assert [(s.point, s.verified, s.use) for s in app.strengths] == [
        ("RAG", True, True), ("Kubernetes", False, False), ("Docker", True, True)]  # first 4, without the repeat
    assert app.plan.candidate_name == "Alex Muster"


def test_plan_prompt_has_the_job_ad_and_cv():
    client = FakeClient(plan())
    agent.make_plan(agent.start(JOB_AD, CV, "message"), client, model="some-model")
    (call,) = client.calls
    assert call["model"] == "some-model" and call["think"] is False and call["format"] == agent.Plan.model_json_schema()
    assert "short recruiter message" in call["messages"][0]["content"]
    assert JOB_AD.strip() in call["messages"][1]["content"] and CV.strip() in call["messages"][1]["content"]


def test_the_writer_only_gets_the_strengths_you_kept():
    app = planned(gap="Kubernetes", gap_answer="Learned Docker on the job")
    app.strengths[1].use = False
    text = agent.plan_for_writer(app)
    assert "1. RAG" in text and "Docker (the job" not in text
    assert "Gap to address honestly: Kubernetes. How: Learned Docker on the job" in text


# ----------------------------------------------------------------------------- writing, checking, fixing
def test_a_good_first_draft_needs_no_fixing():
    app = planned()
    client = FakeClient(draft())
    agent.write_draft(app, client)
    assert len(client.calls) == 1 and len(app.versions) == 1
    assert app.latest.why == "first draft" and app.latest.problems == []
    system, user = (m["content"] for m in client.calls[0]["messages"])
    assert 'Mention the job title "Junior AI Engineer (m/f/d)" and the company "Brightline Analytics GmbH"' in system
    assert user.startswith("APPROVED PLAN:\nCandidate: Alex Muster")


def test_problems_go_back_to_the_agent_until_fixed():
    app = planned()
    client = FakeClient(draft(greeting="Yo team,"), draft())
    agent.write_draft(app, client)
    assert [v.why for v in app.versions] == ["first draft", "fixed: Start with a greeting like 'Dear <contact person>' "
                                             "if one is given, otherwise 'Dear Hiring Team' ('Hello' for a message)."]
    fix_prompt = client.calls[1]["messages"][1]["content"]
    assert "CURRENT DRAFT:\n<<<\nApplication: Junior AI Engineer" in fix_prompt
    assert fix_prompt.endswith("FIX ONLY THESE PROBLEMS, keep everything else:\n- Start with a greeting like "
                               "'Dear <contact person>' if one is given, otherwise 'Dear Hiring Team' ('Hello' for a message).")
    assert app.latest.problems == []


def test_the_ai_never_writes_the_closing():
    assert agent.DraftAnswer.model_json_schema()["required"] == ["subject", "greeting", "paragraphs"]
    app = planned(language="Deutsch", kind="message")
    paragraphs = ["Ich baue RAG Chatbots.", "Hätten Sie Zeit für ein Gespräch?\nMit freundlichen Grüßen\nAlex Muster",
                  "Viele Grüße", "Alex Muster", " "]
    agent.write_draft(app, FakeClient(*[draft(paragraphs=paragraphs, greeting="Guten Tag,")] * TRIES))
    d = app.versions[0].draft
    assert d.paragraphs == ["Ich baue RAG Chatbots.", "Hätten Sie Zeit für ein Gespräch?"]  # nothing twice
    assert d.closing == "Viele Grüße\nAlex Muster"  # added by Python


def test_the_greeting_always_ends_with_a_comma():
    app = planned()
    agent.write_draft(app, FakeClient(draft(greeting="Dear Hiring Team")))
    assert app.latest.draft.greeting == "Dear Hiring Team,"


def test_a_last_paragraph_that_only_starts_like_a_closing_stays():
    app = planned()
    agent.write_draft(app, FakeClient(draft(paragraphs=GOOD_PARAGRAPHS[:-1] + ["Best of all, I could start soon and talk in an interview."])))
    assert app.versions[0].draft.paragraphs[-1] == "Best of all, I could start soon and talk in an interview."


def test_the_fix_loop_stops_after_two_rounds():
    app = planned()
    client = FakeClient(*[draft(greeting="Yo,")] * TRIES)
    agent.write_draft(app, client)
    assert len(app.versions) == 1 + agent.MAX_FIX_ROUNDS and app.latest.problems  # you see what's still wrong


def test_writing_needs_a_plan_and_a_strength():
    app = agent.start(JOB_AD, CV)
    with pytest.raises(agent.AgentError, match="Make a plan first"):
        agent.write_draft(app, FakeClient())
    app = planned()
    for s in app.strengths:
        s.use = False
    with pytest.raises(agent.AgentError, match="at least one strength"):
        agent.write_draft(app, FakeClient())


def test_your_feedback_is_applied_and_its_facts_may_be_used():
    app = planned()
    agent.write_draft(app, FakeClient(draft()))
    new = GOOD_PARAGRAPHS[:-1] + ["I can start on 1 March 2027 and would be glad to talk in an interview."]
    client = FakeClient(draft(paragraphs=new))
    agent.revise(app, "Mention that I can start on 1 March 2027.", client)
    assert "1 March 2027" in app.notes
    assert app.latest.why == "your feedback: Mention that I can start on 1 March 2027." and app.latest.problems == []
    assert client.calls[0]["messages"][1]["content"].endswith(
        "THE CANDIDATE'S FEEDBACK, apply it:\nMention that I can start on 1 March 2027.")


@pytest.mark.parametrize("feedback, limit", [
    ("At most 180 words.", 180), ("Bitte höchstens 120 Wörter.", 120), ("max 150 words please", 150),
    ("A bit shorter.", 180),  # 85% of the 212 words now
    ("Make it shorter, at most 250 words.", 180),  # both: shorter than now wins over the looser limit
    ("Much shorter!", 180), ("More formal, please.", None),
])
def test_word_limits_in_your_feedback_become_checks(feedback, limit):
    app = planned()
    agent.write_draft(app, FakeClient(draft()))
    # each fix round: the AI's cut list (empty here, so nothing goes), then a rewrite
    agent.revise(app, feedback, FakeClient(draft(), *[cut_reply(), draft()] * agent.MAX_FIX_ROUNDS))
    assert app.max_words == limit
    if limit:  # Python checks your limit, and you see that the agent couldn't meet it
        assert app.latest.problems[0].startswith(f"It is too long: 212 words. Write at most {limit}: cut about")
        assert len(app.versions) == 2 + agent.MAX_FIX_ROUNDS


def test_the_writer_is_told_the_same_word_range_python_checks():
    app = planned()
    client = FakeClient(draft(), draft(), *[cut_reply(), draft()] * agent.MAX_FIX_ROUNDS)
    agent.write_draft(app, client)
    assert "3 or 4 paragraphs and 220 to 340 words in total" in client.calls[0]["messages"][0]["content"]
    agent.revise(app, "at most 180 words", client)
    assert "3 or 4 paragraphs and 117 to 162 words in total" in client.calls[1]["messages"][0]["content"]
    assert agent.word_limits(app) == ((108, 180), (117, 162))


def test_too_long_only_is_fixed_by_cutting_sentences_python_counts():
    app = planned()
    agent.write_draft(app, FakeClient(draft()))
    app.notes = "I can start on 1 March 2027."
    long = GOOD_PARAGRAPHS[:-1] + ["I can start on 1 March 2027. I would be glad to talk in an interview."]
    words_before = agent.word_count(agent.Draft.model_validate_json(draft(paragraphs=long)).body)
    assert words_before > 190  # so the limit below makes it too long
    client = FakeClient(draft(paragraphs=long), cut_reply("s1", "s3", "s6", "s7", "s8"))
    agent.revise(app, "at most 190 words", client)
    cut_call = client.calls[1]
    assert cut_call["format"]["properties"]["remove"]["items"]["enum"][:2] == ["s1", "s2"]
    assert f"It has {words_before} words and must have at most 190" in cut_call["messages"][0]["content"]
    final = app.latest
    assert final.why.startswith("cut: It is too long") and final.problems == []
    text = final.draft.body
    assert "Brightline Analytics builds AI assistants" in text  # s1 names the company: protected
    assert "I can start on 1 March 2027." in text and "talk in an interview" in text  # your fact, the final ask
    assert "As a working student I built a RAG chatbot" in text  # stopped cutting once it was short enough
    assert "I connected it to the OpenAI API" not in text  # s3: cut


def test_a_cut_that_removes_nothing_falls_back_to_a_rewrite():
    app = planned()
    agent.write_draft(app, FakeClient(draft()))
    client = FakeClient(draft(), *[cut_reply("s1"), draft()] * agent.MAX_FIX_ROUNDS)  # s1 is protected: nothing goes
    agent.revise(app, "at most 200 words", client)
    assert [v.why[:6] for v in app.versions[2:]] == ["fixed:"] * agent.MAX_FIX_ROUNDS


def test_shorter_never_goes_below_three_quarters_of_the_minimum():
    app = planned()
    agent.write_draft(app, FakeClient(*[draft(paragraphs=GOOD_PARAGRAPHS[:2])] * TRIES))  # a short letter, fixed in vain
    agent.revise(app, "shorter", FakeClient(draft(), *[cut_reply(), draft()] * agent.MAX_FIX_ROUNDS))
    assert app.max_words == 150


def test_empty_feedback_and_nothing_to_revise():
    app = planned()
    with pytest.raises(agent.AgentError, match="no draft"):
        agent.revise(app, "shorter", FakeClient())
    agent.write_draft(app, FakeClient(draft()))
    with pytest.raises(agent.AgentError, match="Write what should change"):
        agent.revise(app, "  ", FakeClient())


def test_broken_answer_is_retried_with_a_little_randomness():
    app = planned()
    client = FakeClient('{"subject": "x"', draft())
    agent.write_draft(app, client)
    assert [c["options"]["temperature"] for c in client.calls] == [0, 0.3]
    assert all(c["stream"] for c in client.calls)


def test_a_model_stuck_on_blank_lines_is_stopped_and_retried():
    app = planned()
    stuck = '{"subject": "x", "greeting": "Dear Hiring Team,", "paragraphs": ["Hi"]' + "\n  " * 400 + "}"
    client = FakeClient(stuck, stuck, draft())
    agent.write_draft(app, client)
    assert [c["options"]["temperature"] for c in client.calls] == [0, 0.3, 0.7]
    assert app.latest.problems == []


def test_three_bad_answers_give_a_clear_error():
    with pytest.raises(agent.AgentError, match="not in the expected format, 3 times"):
        agent.make_plan(agent.start(JOB_AD, CV), FakeClient("{", "nope", '{"candidate_name": 1}'))


# ----------------------------------------------------------------------------- approving and exporting
def test_approve_saves_text_and_history(tmp_path):
    app = planned()
    agent.write_draft(app, FakeClient(draft(greeting="Yo,"), draft()))
    path = agent.approve(app, tmp_path)
    assert app.approved and path.name.endswith("-brightline-analytics-cover_letter.txt")
    assert path.read_text().startswith("Application: Junior AI Engineer at Brightline Analytics\n\nDear Hiring Team,")
    history = json.loads(path.with_suffix(".json").read_text())
    assert [v["why"][:11] for v in history["versions"]] == ["first draft", "fixed: Star"]


def test_nothing_to_approve_yet(tmp_path):
    with pytest.raises(agent.AgentError, match="no draft to approve"):
        agent.approve(planned(), tmp_path)


def test_word_file():
    from docx import Document

    app = planned()
    agent.write_draft(app, FakeClient(draft()))
    paragraphs = [p.text for p in Document(io.BytesIO(agent.to_docx(app))).paragraphs]
    assert paragraphs[:2] == ["Application: Junior AI Engineer at Brightline Analytics", "Dear Hiring Team,"]
    assert paragraphs[-2:] == ["Kind regards", "Alex Muster"]


def test_you_approve_in_the_terminal(monkeypatch, tmp_path, capsys):
    client = FakeClient(plan(), draft(greeting="Yo,"), draft(), draft())
    answers = iter(["2", "Make it warmer.", ""])  # drop strength 2, give feedback once, then approve
    monkeypatch.setattr(agent.ollama, "Client", lambda: client)
    monkeypatch.setattr(agent, "LETTERS_DIR", tmp_path)
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    monkeypatch.setattr(sys, "argv", ["agent.py", str(HERE / "samples/job_ai_engineer_en.txt"), str(HERE / "samples/cv_ai_student.txt")])
    agent.main()
    assert "Docker (the job" not in client.calls[1]["messages"][1]["content"]  # strength 2 was left out
    assert "THE CANDIDATE'S FEEDBACK, apply it:\nMake it warmer." in client.calls[3]["messages"][1]["content"]
    assert "Nothing was sent" in capsys.readouterr().out
    assert len(list(tmp_path.glob("*.txt"))) == 1


# ----------------------------------------------------------------------------- Ollama problems
def test_ollama_not_running():
    with pytest.raises(agent.AgentError, match="Open the Ollama app"):
        agent.make_plan(agent.start(JOB_AD, CV), FakeClient(ConnectionError("refused")))


def test_model_not_downloaded():
    with pytest.raises(agent.AgentError, match="ollama pull gemma4:e4b"):
        agent.make_plan(agent.start(JOB_AD, CV), FakeClient(ollama.ResponseError("not found", 404)))
