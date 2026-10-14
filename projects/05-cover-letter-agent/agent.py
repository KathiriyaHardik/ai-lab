#!/usr/bin/env python3
"""Cover Letter Agent: plans, writes, checks and fixes a cover letter or recruiter message,
and waits for your approval at every important step.

    uv run agent.py samples/job_ai_engineer_en.txt samples/cv_ai_student.txt
    uv run agent.py my-job.txt my-cv.pdf --kind message --language Deutsch

The steps:
  1. PLAN    the AI picks the 3-4 strongest matches between CV and job ad, each with words copied
             from the CV as proof (Python checks every quote), plus one honest gap
  2. YOU     approve the plan: drop strengths, add notes ("I can start in March")
  3. WRITE   the AI writes the text from the approved plan only
  4. CHECK   Python checks the draft: length, company and job title, greeting and closing,
             language, placeholders, and every number must come from the CV, job ad or your notes
  5. FIX     problems go back to the AI, which fixes them by itself (at most 3 rounds). When the only
             problem is length, the AI picks sentences to cut and Python cuts and counts
  6. YOU     approve the draft, or give feedback and the agent revises it (back to 4)

Nothing is ever sent anywhere: you copy or download the approved text and send it yourself.
Everything runs on this computer with Ollama.
"""
import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path

import ollama
from typing import Literal

from pydantic import BaseModel, Field, ValidationError, create_model

MODEL = os.environ.get("AGENT_MODEL", "gemma4:e4b")
LETTERS_DIR = Path(__file__).resolve().parent / "my-letters"  # approved texts (git-ignored)
CONTEXT_TOKENS = 16384
MAX_ANSWER_TOKENS = 2048  # hard stop, so a stuck model can't run forever
MAX_FIX_ROUNDS = 3
STUCK_WHITESPACE = 200  # this much blank space in a row means the model is stuck (see ask)
TEMPERATURES = (0, 0.3, 0.7)  # first try, then a little more randomness for each retry
MAX_INPUT_CHARS = 30000
KINDS = {  # words: what Python accepts. The AI is asked for a narrower range inside it, because it overshoots.
    "cover_letter": dict(name="cover letter", words=(200, 380), target=(220, 340),
                         rules="3 or 4 paragraphs and {low} to {high} words in total: why this job and company, 2 or 3 "
                               "strengths with concrete proof from the CV, and a short last paragraph asking for an interview."),
    "message": dict(name="short recruiter message (LinkedIn or email)", words=(45, 130), target=(60, 110),
                    rules="2 short paragraphs and {low} to {high} words in total: who you are and why this job, your "
                          "strongest match, and a friendly question about a short call."),
}
LANGUAGES = {
    "English": dict(greetings=("dear", "hello", "hi"), sign_off={"cover_letter": "Kind regards", "message": "Best regards"},
                    closings=("kind regards", "best regards", "regards", "sincerely", "yours sincerely", "best wishes"),
                    greeting="'Dear <contact person>' if one is given, otherwise 'Dear Hiring Team' ('Hello' for a message)",
                    closing="'Kind regards'", style="Plain, warm and professional English."),
    "Deutsch": dict(greetings=("sehr geehrte", "guten tag", "hallo"), sign_off={"cover_letter": "Mit freundlichen Grüßen", "message": "Viele Grüße"},
                    closings=("mit freundlichen grüßen", "freundliche grüße", "viele grüße", "beste grüße"),
                    greeting="'Sehr geehrte Frau …' / 'Sehr geehrter Herr …' if a contact person is given, otherwise "
                             "'Sehr geehrte Damen und Herren' ('Guten Tag …' for a message)",
                    closing="'Mit freundlichen Grüßen' ('Viele Grüße' for a message)",
                    style="Correct, natural German. Address the reader formally with 'Sie'."),
}
CLICHES = ("i am writing to apply", "i am writing to express my interest", "hiermit bewerbe ich mich",
           "mit großem interesse habe ich ihre stellenanzeige gelesen")
STOP = {"English": {"the", "and", "i", "my", "to", "of", "with", "for", "in", "is", "a", "your"},
        "Deutsch": {"der", "die", "das", "und", "ich", "mit", "für", "in", "ist", "sie", "ihre", "meine", "zu"}}
LEGAL_FORMS = r"\b(gmbh|mbh|ag|se|kg|ohg|ug|e\s?v|inc|ltd|llc|co)\b\.?|&"
# your feedback "at most 250 words" / "höchstens 120 Wörter" / "shorter" becomes a limit Python checks
LIMIT_RE = re.compile(r"(?:at most|max(?:imum)?|no more than|under|below|höchstens|maximal|unter|weniger als)\s+"
                      r"(\d{2,4})\s*(?:words|wörter|wörtern)\b", re.IGNORECASE)
SHORTER_RE = re.compile(r"\b(shorter|more concise|kürzer|knapper)\b", re.IGNORECASE)

PLAN_PROMPT = """You help a job seeker plan a {kind} for one job. You get the job ad and the CV.

- candidate_name: the candidate's name, as written in the CV.
- company: the company name from the job ad. job_title: the job title from the ad, without (m/f/d).
- contact_person: the person to address if the ad names one (for example "Frau Dr. Weber"), else "".
- strengths: the 3 or 4 strongest matches between the CV and what the job asks for, most important first.
  point: the strength in a few words. requirement: what the job ad asks for. evidence: the exact words from
  the CV that prove it, at most 20 words, copied character for character.
- gap: the most important requirement the CV does not show, or "".
  gap_answer: an honest way to address it with facts from the CV (related experience, learning fast), or "".
- motivation: one sentence on why this job fits the candidate, based on what the job ad says.

Never invent anything that is not in the CV or the job ad."""

WRITE_PROMPT = """You write a {kind} in {language} for a job seeker. Use ONLY the approved plan, the CV, the
job ad and the candidate's notes. Never invent experience, numbers, names, degrees or skills.

Rules:
- {rules}
- Mention the job title "{job_title}" and the company "{company}".
- Greeting: {greeting}. Don't write a closing or the name: they are added automatically.
- Use the approved strengths in their order, each with a concrete fact from the CV. If the plan has a gap,
  address it honestly in one sentence.
- No placeholders like [Company] and no empty phrases like "I am writing to apply" or "hiermit bewerbe ich mich".
- {style}

Fill in: subject (a short subject line), greeting, paragraphs (one string per paragraph, without the
greeting, without a closing and without the name)."""


CUT_PROMPT = """You shorten a {kind} without losing what matters. Its sentences are numbered (s1, s2, ...).
It has {count} words and must have at most {high}, so about {cut} words have to go.

List the sentences to remove, the least important first: repetition, general phrases and weaker points.
Keep the strongest proof from the CV and every fact the candidate asked for."""


class AgentError(Exception):
    """A problem the user can fix (empty input, Ollama not running, model missing...)."""


# ----------------------------------------------------------------------------- the AI's answers
class Strength(BaseModel):
    point: str
    requirement: str
    evidence: str = Field(description="Exact words copied from the CV")


class Plan(BaseModel):
    candidate_name: str
    company: str
    job_title: str
    contact_person: str
    strengths: list[Strength] = Field(min_length=1)
    gap: str
    gap_answer: str
    # One field, not "motivation" + "motivation_evidence": the model treats the two as one idea, stops after
    # the first, and then gets stuck writing blank lines because the shape still wants the second (see README).
    motivation: str


class DraftAnswer(BaseModel):  # what the AI fills in. No closing: when it had to write one, it sometimes put the
    subject: str                 # closing in the paragraphs, then got stuck writing blank lines (see README)
    greeting: str
    paragraphs: list[str] = Field(min_length=1)


class Draft(DraftAnswer):
    closing: str = ""  # added by Python: the sign-off and the candidate's name

    @property
    def body(self):
        return "\n\n".join([self.greeting.strip(), *[p.strip() for p in self.paragraphs if p.strip()], self.closing.strip()])

    @property
    def text(self):
        return f"{self.subject.strip()}\n\n{self.body}"


# ----------------------------------------------------------------------------- the application, step by step
@dataclass
class CheckedStrength:
    point: str
    requirement: str
    evidence: str
    verified: bool  # the evidence really is in the CV
    use: bool = True  # you can switch it off before the agent writes


@dataclass
class Version:
    draft: Draft
    problems: list  # what Python's check still found
    why: str  # "first draft", "fixed: ...", "your feedback: ..."


@dataclass
class Application:
    job_ad: str
    cv: str
    kind: str = "cover_letter"
    language: str = "English"
    notes: str = ""  # your extra facts: start date, availability... (the agent may use them)
    plan: Plan | None = None
    strengths: list = field(default_factory=list)  # CheckedStrength
    address_gap: bool = True  # you can tell the agent to leave the gap out
    max_words: int | None = None  # a limit from your feedback, checked by Python
    versions: list = field(default_factory=list)  # Version, oldest first
    approved: bool = False

    @property
    def latest(self):
        return self.versions[-1] if self.versions else None

    @property
    def sources(self):
        """Everything the letter may take facts from."""
        return "\n".join([self.cv, self.job_ad, self.notes])


# ----------------------------------------------------------------------------- pure checks (no AI)
def words(text):
    """Lower-case words only, so line breaks, quotes and dashes don't block a match."""
    return " ".join(re.findall(r"\w+", unicodedata.normalize("NFKC", text).lower()))


def quote_in(quote, text):
    """True if every part of the quote (split at "...") appears word for word in the text."""
    parts = [p for p in (words(p) for p in re.split(r"\.\.\.|…", quote)) if p]
    haystack = f" {words(text)} "
    return bool(parts) and all(f" {p} " in haystack for p in parts)


def core_name(name):
    """'Brightline Analytics GmbH' -> 'brightline analytics'; 'Junior AI Engineer (m/f/d)' -> 'junior ai engineer'."""
    name = re.sub(r"\(.*?\)", " ", name.lower())
    return words(re.sub(LEGAL_FORMS, " ", name))


def word_limits(app):
    """(lowest, highest) words Python accepts, and (lowest, highest) the AI is asked for. Your limit wins."""
    (low, high), (aim_low, aim_high) = KINDS[app.kind]["words"], KINDS[app.kind]["target"]
    if app.max_words:  # the minimum moves down with your limit, and the AI aims a bit below it
        high, low = app.max_words, min(low, int(app.max_words * 0.6))
        aim_high, aim_low = int(app.max_words * 0.9), min(aim_low, int(app.max_words * 0.65))
    return (low, high), (aim_low, aim_high)


def word_count(text):
    return len(re.findall(r"\w+", text))


def numbers(text):
    """Every number, digits only: '2,000' and '2.000' both become '2000'."""
    return {re.sub(r"\D", "", n) for n in re.findall(r"\d+(?:[.,]\d+)*", text)}


def check_draft(app, draft):
    """Everything Python can check about a draft, as plain-language problems the agent can fix."""
    problems, plan, language = [], app.plan, LANGUAGES[app.language]
    text, low = draft.body, words(draft.body)
    count = word_count(text)
    (low_words, high_words), _ = word_limits(app)
    if count < low_words:
        problems.append(f"It is too short: {count} words. Write at least {low_words}.")
    if count > high_words:
        problems.append(f"It is too long: {count} words. Write at most {high_words}: cut about "
                        f"{count - high_words + 10} words, for example by removing a sentence that repeats something.")
    if core_name(plan.company) and f" {core_name(plan.company)} " not in f" {words(draft.text)} ":
        problems.append(f'Mention the company "{plan.company}".')
    if core_name(plan.job_title) and f" {core_name(plan.job_title)} " not in f" {words(draft.text)} ":
        problems.append(f'Mention the job title "{plan.job_title}".')
    if not words(draft.greeting).startswith(language["greetings"]):
        problems.append(f"Start with a greeting like {language['greeting']}.")
    if not any(c in words(draft.closing) for c in (words(c) for c in language["closings"])):
        problems.append(f"End with {language['closing']}.")
    if words(plan.candidate_name) and words(plan.candidate_name) not in words(draft.closing):
        problems.append(f'Sign with the name "{plan.candidate_name}".')
    placeholder = re.search(r"\[[^\]]*\]|\{[^}]*\}|<[^>]+>|\bX{3,}\b|\.{4,}", draft.text)
    if placeholder:
        problems.append(f'Replace the placeholder "{placeholder.group(0)}" with real content.')
    invented = sorted(numbers(draft.text) - numbers(app.sources), key=int)
    if invented:
        problems.append(f"These numbers are not in the CV, the job ad or the notes: {', '.join(invented)}. "
                        "Remove them or use the real numbers.")
    for cliche in CLICHES:
        if cliche in low:
            problems.append(f'Don\'t use the empty phrase "{cliche}". Start with something concrete.')
    tokens = low.split()
    other = "Deutsch" if app.language == "English" else "English"
    if sum(t in STOP[other] for t in tokens) > sum(t in STOP[app.language] for t in tokens):
        problems.append(f"Write the whole text in {app.language}.")
    return problems


# ----------------------------------------------------------------------------- talking to Ollama
def ollama_call(model, call):
    """Run one Ollama request and turn its errors into messages the user can act on."""
    try:
        return call()
    except ConnectionError as e:
        raise AgentError("Ollama is not running. Open the Ollama app, then try again.") from e
    except ollama.ResponseError as e:
        if e.status_code == 404:
            raise AgentError(f"The AI model {model} is not downloaded yet. Run:  ollama pull {model}") from e
        if "structured output" in str(e.error):
            raise AgentError(
                "This Ollama install can't force the answer into JSON. Install the official app "
                "from ollama.com (Homebrew: brew install --cask ollama-app)."
            ) from e
        raise AgentError(f"Ollama returned an error: {e.error}") from e


def stream(client, model, messages, shape, temperature):
    """The reply, read as it arrives. None if the model gets stuck writing blank space: stop at once, don't wait."""
    text = ""
    chunks = ollama_call(model, lambda: client.chat(
        model=model,
        messages=messages,
        format=shape.model_json_schema(),  # Ollama forces the reply into this JSON shape
        think=False,
        stream=True,
        options={"temperature": temperature, "num_ctx": CONTEXT_TOKENS, "num_predict": MAX_ANSWER_TOKENS},
    ))
    for chunk in chunks:
        text += chunk.message.content or ""
        if len(text) - len(text.rstrip()) > STUCK_WHITESPACE:
            return None
    return text


def ask(client, model, system, user, shape):
    """One AI call whose reply is forced into the JSON shape of the pydantic model `shape`."""
    messages = [{"role": "system", "content": system}, {"role": "user", "content": user}]
    problem = None
    for temperature in TEMPERATURES:  # a broken answer at temperature 0 would break the same way again
        text = ollama_call(model, lambda: stream(client, model, messages, shape, temperature))
        if text is None:
            problem = "the model got stuck writing blank lines"
            continue
        try:
            return shape.model_validate_json(text)
        except ValidationError as e:
            problem = e
    raise AgentError(f"The AI's answer was not in the expected format, {len(TEMPERATURES)} times. Try again.\n\n{problem}")


# ----------------------------------------------------------------------------- the agent's steps
def start(job_ad, cv, kind="cover_letter", language="English", notes=""):
    job_ad, cv = (job_ad or "").strip(), (cv or "").strip()
    if len(job_ad) < 100:
        raise AgentError("The job ad is empty or very short. Paste the full job ad.")
    if len(cv) < 100:
        raise AgentError("The CV is empty or very short. Upload your CV or paste it as text.")
    if len(job_ad) + len(cv) + len(notes) > MAX_INPUT_CHARS:
        raise AgentError("The job ad and CV are too long together. Remove parts that don't matter.")
    if kind not in KINDS or language not in LANGUAGES:
        raise AgentError(f"Unknown kind or language: {kind}, {language}.")
    return Application(job_ad, cv, kind, language, notes.strip())


def make_plan(app, client=None, model=MODEL):
    """Step 1: the AI's plan. Every strength's proof is checked against the CV."""
    client = client or ollama.Client()
    plan = ask(client, model, PLAN_PROMPT.format(kind=KINDS[app.kind]["name"]),
               f"JOB AD:\n<<<\n{app.job_ad}\n>>>\n\nCV:\n<<<\n{app.cv}\n>>>", Plan)
    if plan.candidate_name.isupper():  # CV headings are often in capitals: "ALEX MUSTER" -> "Alex Muster"
        plan.candidate_name = plan.candidate_name.title()
    app.plan = plan
    app.strengths, seen = [], set()
    for s in plan.strengths[:4]:
        if words(s.point) in seen:
            continue
        seen.add(words(s.point))
        verified = quote_in(s.evidence, app.cv)
        app.strengths.append(CheckedStrength(s.point.strip(), s.requirement.strip(), s.evidence.strip(), verified, use=verified))
    return app


def plan_for_writer(app):
    p, chosen = app.plan, [s for s in app.strengths if s.use]
    lines = [f"Candidate: {p.candidate_name}", f"Job: {p.job_title} at {p.company}",
             f"Contact person: {p.contact_person or 'none'}", "Strengths to use, in this order:"]
    lines += [f"{n}. {s.point} (the job asks for: {s.requirement}; proof from the CV: \"{s.evidence}\")"
              for n, s in enumerate(chosen, 1)]
    if p.gap and p.gap_answer and app.address_gap:
        lines.append(f"Gap to address honestly: {p.gap}. How: {p.gap_answer}")
    if p.motivation:
        lines.append(f"Why this job: {p.motivation}")
    return "\n".join(lines)


def write(app, client, model, instruction=""):
    """One call to the writer. `instruction` carries the problems to fix or your feedback."""
    p, kind, language = app.plan, KINDS[app.kind], LANGUAGES[app.language]
    _, (aim_low, aim_high) = word_limits(app)  # the same numbers Python checks, so the two never contradict
    rules = kind["rules"].format(low=aim_low, high=aim_high)
    system = WRITE_PROMPT.format(kind=kind["name"], language=app.language, rules=rules, job_title=p.job_title,
                                 company=p.company, greeting=language["greeting"], style=language["style"])
    user = (f"APPROVED PLAN:\n{plan_for_writer(app)}\n\nCANDIDATE'S NOTES:\n{app.notes or '(none)'}\n\n"
            f"CV:\n<<<\n{app.cv}\n>>>\n\nJOB AD:\n<<<\n{app.job_ad}\n>>>")
    if app.latest:
        user += f"\n\nCURRENT DRAFT:\n<<<\n{app.latest.draft.text}\n>>>"
    if instruction:
        user += f"\n\n{instruction}"
    return finish(app, ask(client, model, system, user, DraftAnswer))


def finish(app, answer):
    """Python adds the closing. A closing or name the AI wrote at the end anyway is removed, so it isn't there twice."""
    language, name = LANGUAGES[app.language], app.plan.candidate_name.strip()
    closing_line = lambda line: (words(line) == words(name) or  # noqa: E731
                                 any(words(line).startswith(c) and len(words(line).split()) <= len(c.split()) + 3
                                     for c in language["closings"]))
    paragraphs = [para.strip() for para in answer.paragraphs if para.strip()]
    while paragraphs:
        lines = paragraphs[-1].splitlines()
        while lines and (not lines[-1].strip() or closing_line(lines[-1])):
            lines.pop()
        if lines:
            paragraphs[-1] = "\n".join(lines).strip()
            break
        paragraphs.pop()
    closing = f"{language['sign_off'][app.kind]}\n{name}"
    greeting = answer.greeting.strip().rstrip(" ,:;!") + ","  # "Dear Hiring Team," / "Sehr geehrte Damen und Herren,"
    return Draft(subject=answer.subject, greeting=greeting, paragraphs=paragraphs or [""], closing=closing)


def sentences(paragraph):
    return [s for s in re.split(r"(?<=[.!?])\s+(?=[A-ZÄÖÜ\"„“])", paragraph.strip()) if s]


def cut(app, draft, client, model):
    """Too long only: the AI picks the sentences to cut, Python cuts them and counts. The AI can't count words
    reliably while it writes, so a rewrite "with at most 183 words" kept landing a few words over."""
    (_, high), _ = word_limits(app)
    parts = [sentences(p) for p in draft.paragraphs]
    places = [(pi, si) for pi, ss in enumerate(parts) for si in range(len(ss))]  # (paragraph, sentence)
    sid_at = {place: f"s{n}" for n, place in enumerate(places, 1)}
    text_of = {sid_at[(pi, si)]: parts[pi][si] for pi, si in places}
    # never cut: the company or job title, a number from your notes (like a start date), the final ask for a talk
    keep_words = [w for w in (core_name(app.plan.company), core_name(app.plan.job_title)) if w]
    notes_numbers = numbers(app.notes) - numbers(app.cv) - numbers(app.job_ad)
    protected = {sid for sid, text in text_of.items()
                 if any(f" {w} " in f" {words(text)} " for w in keep_words) or numbers(text) & notes_numbers}
    protected.add(sid_at[places[-1]])
    count = word_count(draft.body)
    shape = create_model("Cut", remove=(list[Literal[tuple(text_of)]], ...))  # only ids that exist
    numbered = "\n".join(f"{sid}: {text}" for sid, text in text_of.items())
    answer = ask(client, model, CUT_PROMPT.format(kind=KINDS[app.kind]["name"], count=count, high=high, cut=count - high + 10),
                 f"SENTENCES:\n{numbered}\n\nCANDIDATE'S NOTES:\n{app.notes or '(none)'}", shape)
    removed = set()
    for sid in answer.remove:
        if count <= high:
            break
        if sid not in protected and sid not in removed:
            removed.add(sid)
            count -= word_count(text_of[sid])
    paragraphs = [" ".join(t for si, t in enumerate(ss) if sid_at[(pi, si)] not in removed) for pi, ss in enumerate(parts)]
    return Draft(subject=draft.subject, greeting=draft.greeting, paragraphs=[p for p in paragraphs if p.strip()] or [""],
                 closing=draft.closing)


def check_and_fix(app, draft, why, client, model):
    """Python checks the draft; the agent fixes what it finds, at most MAX_FIX_ROUNDS times."""
    problems = check_draft(app, draft)
    app.versions.append(Version(draft, problems, why))
    for _ in range(MAX_FIX_ROUNDS):
        if not problems:
            break
        shorter = cut(app, draft, client, model) if all(p.startswith("It is too long") for p in problems) else None
        if shorter and shorter.body != draft.body:
            draft, label = shorter, "cut: "
        else:  # other problems, or the cut removed nothing: rewrite
            fix = "FIX ONLY THESE PROBLEMS, keep everything else:\n" + "\n".join(f"- {p}" for p in problems)
            draft, label = write(app, client, model, fix), "fixed: "
        problems = check_draft(app, draft)
        app.versions.append(Version(draft, problems, label + "; ".join(app.latest.problems)))
    return app


def write_draft(app, client=None, model=MODEL):
    """Step 3-5: write from the approved plan, then check and fix."""
    if not app.plan:
        raise AgentError("Make a plan first.")
    if not any(s.use for s in app.strengths):
        raise AgentError("Choose at least one strength for the agent to write about.")
    client = client or ollama.Client()
    return check_and_fix(app, write(app, client, model), "first draft", client, model)


def revise(app, feedback, client=None, model=MODEL):
    """Step 6: your feedback. New facts in it ("I can start in March 2027") become notes the agent may use."""
    feedback = (feedback or "").strip()
    if not feedback:
        raise AgentError("Write what should change, or approve the draft.")
    if not app.latest:
        raise AgentError("There is no draft to revise yet.")
    client = client or ollama.Client()
    # The AI can't count words reliably, so Python turns "at most 250 words" and "shorter" into a limit it checks,
    # and the fix loop enforces it. "Shorter, at most 250" means both: shorter than now, and not over 250.
    limits = [int(n) for n in LIMIT_RE.findall(feedback)]
    if SHORTER_RE.search(feedback):
        limits.append(max(int(KINDS[app.kind]["words"][0] * 0.75), int(word_count(app.latest.draft.body) * 0.85)))
    if limits:
        app.max_words = min(limits)
    app.notes = f"{app.notes}\n{feedback}".strip()
    draft = write(app, client, model, f"THE CANDIDATE'S FEEDBACK, apply it:\n{feedback}")
    return check_and_fix(app, draft, f"your feedback: {feedback}", client, model)


def approve(app, folder=None):
    """You approve the latest draft. It's saved in my-letters/ with its history; nothing is sent."""
    if not app.latest:
        raise AgentError("There is no draft to approve yet.")
    app.approved = True
    folder = Path(folder or LETTERS_DIR)
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"{dt.date.today():%Y-%m-%d}-{slug(app.plan.company)}-{app.kind}"
    (folder / f"{stem}.txt").write_text(app.latest.draft.text + "\n")
    history = {"job_title": app.plan.job_title, "company": app.plan.company, "kind": app.kind, "language": app.language,
               "notes": app.notes, "plan": plan_for_writer(app),
               "versions": [{"why": v.why, "problems": v.problems, "text": v.draft.text} for v in app.versions]}
    (folder / f"{stem}.json").write_text(json.dumps(history, ensure_ascii=False, indent=2))
    return folder / f"{stem}.txt"


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", core_name(text) or "job").strip("-")[:40]


def to_docx(app):
    """The approved text as a Word file."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    draft = app.latest.draft
    doc.add_paragraph().add_run(draft.subject.strip()).bold = True
    for part in [draft.greeting, *draft.paragraphs]:
        if part.strip():
            doc.add_paragraph(part.strip())
    for line in draft.closing.strip().splitlines():
        doc.add_paragraph(line.strip())
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def read_file(path):
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return path.read_text(errors="replace")


# ----------------------------------------------------------------------------- command line (you approve in the terminal)
def show_plan(app):
    print(f"\nPLAN for {app.plan.job_title} at {app.plan.company}\n")
    for n, s in enumerate(app.strengths, 1):
        mark = "✅" if s.verified else "⚠️  proof not found in your CV, switched off"
        print(f"  {n}. {s.point}  ({mark})\n     job asks for: {s.requirement}\n     your CV: \"{s.evidence}\"")
    if app.plan.gap:
        print(f"\n  Gap: {app.plan.gap}\n  How to address it: {app.plan.gap_answer}")


def show_draft(app):
    v = app.latest
    print(f"\n{'-' * 70}\n{v.draft.text}\n{'-' * 70}")
    print(f"({len(app.versions)} version(s) so far; " + (f"still {len(v.problems)} problem(s): " + " | ".join(v.problems)
                                                         if v.problems else "all checks passed") + ")")


def main():
    parser = argparse.ArgumentParser(description="Plan, write and check a cover letter with you in the loop.")
    parser.add_argument("job_ad", help="job ad as a .txt file")
    parser.add_argument("cv", help="CV as a .pdf, .txt or .md file")
    parser.add_argument("--kind", default="cover_letter", choices=list(KINDS))
    parser.add_argument("--language", default="English", choices=list(LANGUAGES))
    parser.add_argument("--notes", default="", help='extra facts the agent may use, e.g. "I can start in March 2027"')
    parser.add_argument("--model", default=MODEL)
    args = parser.parse_args()
    try:
        app = start(read_file(args.job_ad), read_file(args.cv), args.kind, args.language, args.notes)
        print("Planning…")
        make_plan(app, model=args.model)
        show_plan(app)
        drop = input("\nPress Enter to approve the plan, or type the numbers of strengths to leave out (e.g. 3 4): ")
        for n in re.findall(r"\d+", drop):
            if 1 <= int(n) <= len(app.strengths):
                app.strengths[int(n) - 1].use = False
        print("Writing and checking…")
        write_draft(app, model=args.model)
        while True:
            show_draft(app)
            feedback = input("\nPress Enter to approve, or type what should change: ").strip()
            if not feedback:
                break
            print("Revising and checking…")
            revise(app, feedback, model=args.model)
        print(f"\nApproved and saved to {approve(app)}. Nothing was sent: copy it into your application yourself.")
    except AgentError as e:
        sys.exit(f"Error: {e}")
    except (KeyboardInterrupt, EOFError):
        sys.exit("\nStopped. Nothing was saved or sent.")


if __name__ == "__main__":
    main()
