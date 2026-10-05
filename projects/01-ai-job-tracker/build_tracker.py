#!/usr/bin/env python3
"""Build the AI Job Application CRM workbook (.xlsx, works in Excel 2010+ and Google Sheets).

    python build_tracker.py                          # clean tracker -> AI_Job_Application_CRM.xlsx
    python build_tracker.py --sample-data test.xlsx  # same workbook filled with dummy rows (testing only)

Portability rules this file sticks to:
  * no macros, no Excel Tables, no dynamic-array functions (FILTER/SORT/UNIQUE/XLOOKUP)
  * conditional formatting only references cells on its own sheet (Google Sheets requirement),
    so each data sheet keeps a hidden local copy of the thresholds it needs
  * every threshold and dropdown list lives on the Settings sheet; formulas reach the data
    through named ranges (App_*, Ct_*, Int_*, Set_*), so extending a range is a one-place edit
"""
import argparse
import datetime as dt
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as col_idx, get_column_letter as col_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = Path(__file__).with_name("AI_Job_Application_CRM.xlsx")

# ----------------------------------------------------------------------------- design tokens
FONT = "Arial"
INK, TEXT, MUTED, FAINT = "0F172A", "1E293B", "64748B", "94A3B8"
LINE, LINE_SOFT, WHITE = "E2E8F0", "F1F5F9", "FFFFFF"
HEAD_BG, HEAD_AUTO_BG = "1E293B", "64748B"
AUTO_BG, CARD_BG, ZEBRA = "F1F5F9", "F8FAFC", "FAFBFD"
ACCENT, ACCENT_SOFT = "2563EB", "60A5FA"
INPUT_BG, INPUT_FG = "FEF9C3", "1D4ED8"
RED_BG, RED_FG = "FEE2E2", "991B1B"
AMBER_BG, AMBER_FG = "FEF3C7", "92400E"
YELLOW_BG, YELLOW_FG = "FEF9C3", "854D0E"
GREEN_FG = "166534"

DATE_FMT, DAY_FMT, ID_FMT, PCT_FMT = "dd.mm.yyyy", "ddd dd.mm.yyyy", '"APP-"000', "0%"

# status -> (cell fill, text colour, bar colour). Wishlist is deliberately neutral.
STATUS_STYLE = {
    "Wishlist": (None, "475569", "94A3B8"),
    "Applied": ("DBEAFE", "1E40AF", "3B82F6"),
    "Screening": ("FEF3C7", "92400E", "F59E0B"),
    "Interview 1": ("FFEDD5", "9A3412", "F97316"),
    "Interview 2": ("FFEDD5", "9A3412", "F97316"),
    "Final": ("EDE9FE", "5B21B6", "8B5CF6"),
    "Offer": ("DCFCE7", "166534", "22C55E"),
    "Accepted": ("15803D", "FFFFFF", "15803D"),
    "Rejected": ("FEE2E2", "991B1B", "EF4444"),
    "Withdrawn": ("E5E7EB", "4B5563", "9CA3AF"),
}
# (category, lower bound, label, fill, text)
MATCH_BANDS = [
    ("Excellent Match", 90, "90–100%", "DCFCE7", "166534"),
    ("Strong Match", 75, "75–89%", "ECFCCB", "3F6212"),
    ("Moderate Match", 60, "60–74%", "FEF3C7", "92400E"),
    ("Low Match", 0, "below 60%", "F3F4F6", "6B7280"),
]
# action flags in urgency order: (flag, fill, text, meaning)
FLAGS = [
    ("FOLLOW UP REQUIRED", "FEE2E2", "991B1B", "Next Action Date is in the past and the application is still open."),
    ("ACTION DUE TODAY", "FDE68A", "78350F", "Next Action Date is today."),
    ("NO RESPONSE", "E0E7FF", "3730A3",
     "Status is Applied, no First Response Date, no Next Action Date set, and Date Applied is at least "
     "'Default Follow-up Days' ago (Settings)."),
    ("DUE SOON", "FEF9C3", "854D0E", "Next Action Date falls within the Due-Soon Window (Settings)."),
]
PIPELINE = ["Wishlist", "Applied", "Screening", "Interview 1", "Interview 2", "Final", "Offer", "Accepted"]
CLOSED = ["Rejected", "Withdrawn", "Accepted"]

# ----------------------------------------------------------------------------- lists & settings
LISTS = {
    "Job Type": ["Full-time", "Werkstudent", "Internship", "Working Student", "Contract", "Freelance"],
    "Work Mode": ["Remote", "Hybrid", "On-site"],
    "Source": ["LinkedIn", "Company Website", "Indeed", "StepStone", "Glassdoor", "Recruiter", "Referral",
               "University", "Networking", "Other"],
    "Priority": ["High", "Medium", "Low"],
    "Status": ["Wishlist", "Applied", "Screening", "Interview 1", "Interview 2", "Final", "Offer", "Accepted",
               "Rejected", "Withdrawn"],
    "Cover Letter": ["Yes", "No", "Not Required"],
    "Yes / No": ["Yes", "No"],
    "CV Version": ["AI Engineer CV", "Generative AI CV", "AI Automation CV", "AI Solutions CV",
                   "General Software CV", "Werkstudent AI CV", "Custom"],
    "Role Category": ["AI Engineer", "Generative AI Engineer", "Agentic AI Engineer", "AI Automation Engineer",
                      "AI Solutions Engineer", "AI/ML Engineer", "Applied AI Engineer", "AI Consultant",
                      "AI Product / Solutions", "Werkstudent AI", "Other"],
    "Stage When Closed": ["Applied", "Screening", "Interview 1", "Interview 2", "Final", "Offer"],
    "Industry": ["AI / ML", "SaaS", "Enterprise Software", "Automotive", "Consulting", "Data", "Automation",
                 "FinTech", "HealthTech", "B2B SaaS", "Startup", "Other"],
    "AI Relevance": ["High", "Medium", "Low"],
    "Company Status": ["Researching", "Watching", "Networking", "Ready to Apply", "Applied", "Interviewing",
                       "Offer", "Not a Fit"],
    "Relationship": ["Recruiter", "Hiring Manager", "Employee", "Alumni", "Former Colleague",
                     "Friend / Personal", "Event Contact", "Other"],
    "Response": ["No Response", "Positive", "Neutral", "Negative", "Meeting", "Referral"],
    "Referral Potential": ["High", "Medium", "Low", "Unknown"],
    "Interview Stage": ["Screening", "Interview 1", "Interview 2", "Interview 3+", "Final", "Other"],
    "Interview Type": ["Recruiter", "HR", "Technical", "Coding", "AI/ML Technical", "Hiring Manager",
                       "Case Study", "Final", "Other"],
    "Preparation Status": ["Not Started", "In Progress", "Ready", "Completed"],
    "My Performance": ["Poor", "Needs Improvement", "Good", "Very Good", "Excellent"],
    "Skill Tags": ["LLMs", "RAG", "Agentic AI", "Multi-Agent Systems", "Python", "LLM APIs", "Vector Databases",
                   "SQL", "PostgreSQL", "Machine Learning", "AI Automation", "Knowledge Management",
                   "Human-in-the-Loop", "AI Agents"],
}
FIXED_LISTS = {"Status", "Stage When Closed"}  # formulas match these names literally
LIST_HEAD_ROW, LIST_FIRST_ROW, LIST_FIRST_COL = 12, 13, 2

# (named range, label, default, min, max, explanation) -> Settings!E5:E8
SETTINGS = [
    ("Set_FollowUpDays", "Default Follow-up Days", 7, 1, 90,
     "Days after Date Applied with no reply (and no Next Action Date set) before an application is flagged "
     "NO RESPONSE. This is your own rule of thumb, not a fact about any company."),
    ("Set_DueSoonDays", "Due-Soon Window (days)", 3, 0, 30,
     "Next Action, Follow-up and Next Step dates within this many days turn yellow and show as DUE SOON."),
    ("Set_InterviewWindow", "Upcoming Interview Window (days)", 14, 1, 90,
     "Interviews within this many days are listed under Upcoming Interviews on the Follow-Ups tab."),
    ("Set_MinSample", "Minimum Applications for a 'Best' Callout", 5, 1, 100,
     "Analytics only names a CV version, source, role or match band as 'best' once that group has at "
     "least this many applications, so one lucky application doesn't mislead you."),
]
SETTINGS_ROW0 = 5

# ----------------------------------------------------------------------------- sheet geometry
HDR, FIRST = 4, 5  # header row and first data row on every data sheet
ROWS = {"Applications": 1000, "Interviews": 300, "Companies": 300, "Contacts": 500}


def last_row(sheet):
    return FIRST + ROWS[sheet] - 1


# ----------------------------------------------------------------------------- style helpers
def F(size=10, bold=False, color=TEXT, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def put(ws, ref, value=None, font=None, bg=None, align=None, fmt=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if bg is not None:
        c.fill = fill(bg)
    if align is not None:
        c.alignment = align
    if fmt is not None:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c


def cf(ws, rng, formula, bg=None, fg=None, bold=None, stop=True):
    """Formula-based conditional format. The formula is written for the top-left cell of `rng`."""
    font = Font(color=fg, bold=bold) if (fg or bold) else None
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[formula], fill=fill(bg) if bg else None, font=font, stopIfTrue=stop))


def status_cf(ws, rng, first_cell):
    for status, (bg, fg, _) in STATUS_STYLE.items():
        cf(ws, rng, f'{first_cell}="{status}"', bg, fg, bold=status in ("Accepted", "Offer"))


def flag_cf(ws, rng, first_cell):
    for flag, bg, fg, _ in FLAGS:
        cf(ws, rng, f'{first_cell}="{flag}"', bg, fg, bold=True)


def text_colour_cf(ws, rng, first_cell, mapping):
    for value, (fg, bold) in mapping.items():
        cf(ws, rng, f'{first_cell}="{value}"', fg=fg, bold=bold)


def date_due_cf(ws, rng, cell, window_cell, open_condition=None):
    """Red when overdue, yellow when due within the window. `cell` like $O5."""
    extra = f",{open_condition}" if open_condition else ""
    cf(ws, rng, f"AND(ISNUMBER({cell}),INT({cell})<TODAY(){extra})", RED_BG, RED_FG, bold=True)
    cf(ws, rng, f"AND(ISNUMBER({cell}),INT({cell})>=TODAY(),INT({cell})-TODAY()<={window_cell}{extra})",
       YELLOW_BG, YELLOW_FG)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def chart_text(size=800, color=MUTED, bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


# ----------------------------------------------------------------------------- data-sheet builder
class Col:
    """One column on a data sheet.

    kind: text | wrap | date | id | match | int | list | url | auto | helper
    """

    def __init__(self, header, width, kind="text", source=None, formula=None, name=None, prompt=None,
                 align=None, fmt=None):
        self.header, self.width, self.kind = header, width, kind
        self.source, self.formula, self.name, self.prompt = source, formula, name, prompt
        self.align, self.fmt = align, fmt

    @property
    def is_auto(self):
        return self.kind in ("auto", "helper")


def list_ref(name):
    i = list(LISTS).index(name)
    col = col_letter(LIST_FIRST_COL + i)
    return f"Settings!${col}${LIST_FIRST_ROW}:${col}${LIST_FIRST_ROW + len(LISTS[name]) - 1}"


def list_cell(name, i):
    col = col_letter(LIST_FIRST_COL + list(LISTS).index(name))
    return f"Settings!${col}${LIST_FIRST_ROW + i}"


def build_data_sheet(wb, ws, cols, title, subtitle, hint, freeze, prefix, groups=()):
    sheet = ws.title
    last = last_row(sheet)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    put(ws, "A1", title, F(16, True, INK), align=Alignment(vertical="center"))
    put(ws, "A2", subtitle, F(9, color=MUTED))
    if hint:
        put(ws, hint[0], hint[1], F(8, color=FAINT, italic=True))
    for ref, label, color in groups:
        put(ws, ref, label, F(8, True, color))
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[HDR].height = 34

    head_align = Alignment(vertical="center", wrap_text=True, indent=1)
    border = Border(bottom=side(LINE), right=side(LINE_SOFT))
    for i, col in enumerate(cols, 1):
        L = col_letter(i)
        ws.column_dimensions[L].width = col.width
        if col.kind == "helper":
            ws.column_dimensions[L].hidden = True
        put(ws, f"{L}{HDR}", col.header, F(9, True, WHITE), HEAD_AUTO_BG if col.is_auto else HEAD_BG,
            head_align, border=Border(right=side(WHITE)))

        centred = col.kind in ("date", "id", "match", "int") or col.align == "center"
        align = Alignment(vertical="top", wrap_text=col.kind == "wrap",
                          horizontal="center" if centred else "left", indent=0 if centred else 1)
        font = F(10, color=MUTED if col.is_auto else TEXT, underline="single" if col.kind == "url" else None)
        if col.kind == "url":
            font = F(10, color=ACCENT)
        fmt = col.fmt or {"date": DATE_FMT, "id": ID_FMT, "match": '0"%"', "int": "0"}.get(col.kind)
        bg = fill(AUTO_BG) if col.is_auto else None
        for r in range(FIRST, last + 1):
            c = ws.cell(r, i)
            if col.formula:
                c.value = col.formula.replace("{r}", str(r))  # not .format(): array constants use braces
            c.font, c.alignment, c.border = font, align, border
            if fmt:
                c.number_format = fmt
            if bg:
                c.fill = bg

        rng = f"{L}{FIRST}:{L}{last}"
        if col.name:
            add_name(wb, f"{prefix}_{col.name}", f"'{sheet}'!${L}${FIRST}:${L}${last}")
        dv = None
        if col.kind == "list":
            dv = DataValidation(type="list", formula1=list_ref(col.source), allow_blank=True)
            dv.errorStyle = "warning" if col.source == "Relationship" else "stop"
            dv.errorTitle, dv.error = "Pick from the list", f"Choose a value from the dropdown (list: Settings → {col.source})."
        elif col.kind == "date":
            dv = DataValidation(type="date", operator="between", formula1="43831", formula2="49674", allow_blank=True)
            dv.errorTitle, dv.error = "Date needed", "Enter a real date, e.g. 05.10.2026 (Ctrl+; inserts today)."
        elif col.kind == "match":
            dv = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
            dv.errorTitle, dv.error = "0–100", "Enter a whole number from 0 to 100, e.g. 85 (no % sign)."
        elif col.kind == "id":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True)
            dv.errorTitle, dv.error = "Whole number", "Enter a whole number, e.g. 12 (it displays as APP-012). See 'Next ID' above the table."
        elif col.prompt:
            dv = DataValidation(type="custom", formula1="TRUE", allow_blank=True)  # prompt-only
        if dv is not None:
            dv.showErrorMessage = True
            if col.prompt:
                dv.promptTitle, dv.prompt, dv.showInputMessage = col.header[:32], col.prompt, True
            ws.add_data_validation(dv)
            dv.add(rng)

    n = len(cols)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{HDR}:{col_letter(n)}{last}"
    ws.print_title_rows = f"{HDR}:{HDR}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return last


# ----------------------------------------------------------------------------- Applications
def applications_columns():
    closed = '$M{r}="Rejected",$M{r}="Withdrawn"'
    stages = ",".join(f'"{s}"' for s in PIPELINE)
    flags = ",".join(f'"{f[0]}"' for f in FLAGS)
    return [
        Col("Application ID", 10, "id", name="ID"),
        Col("Date Added", 12, "date"),
        Col("Date Applied", 12, "date", name="DateApplied"),
        Col("Company", 22, name="Company"),
        Col("Job Title", 30, name="JobTitle"),
        Col("Job Type", 15, "list", "Job Type"),
        Col("Location", 16),
        Col("Remote / Hybrid / On-site", 15, "list", "Work Mode"),
        Col("Job URL", 24, "url", name="URL"),
        Col("Source", 16, "list", "Source", name="Source"),
        Col("Priority", 10, "list", "Priority", name="Priority", align="center"),
        Col("Match %", 9, "match", name="Match",
            prompt="Your personal fit estimate, 0–100 (type 85, not 85%). It's a judgement call, not an "
                   "objective score. 90+ Excellent · 75–89 Strong · 60–74 Moderate · below 60 Low."),
        Col("Status", 13, "list", "Status", name="Status"),
        Col("Next Action", 28, "wrap", name="NextAction"),
        Col("Next Action Date", 12, "date", name="NextDate"),
        Col("Recruiter / Hiring Manager", 20, name="Contact"),
        Col("Recruiter LinkedIn", 22, "url"),
        Col("CV Version", 18, "list", "CV Version", name="CV"),
        Col("Cover Letter", 12, "list", "Cover Letter", align="center"),
        Col("Referral", 10, "list", "Yes / No", align="center"),
        Col("Salary / Salary Range", 15),
        Col("Key Technologies / Skills", 30, "wrap", name="Skills",
            prompt="Several tags, separated by commas, e.g. LLMs, RAG, Python, Vector Databases. "
                   "Tags listed on Settings (Skill Tags) are counted in Analytics."),
        Col("Why Me?", 46, "wrap",
            prompt="1–2 sentences on why you fit this role. Reuse it for cover letters, recruiter messages "
                   "and 'why are you a good fit?' answers."),
        Col("Notes", 36, "wrap"),
        Col("Last Updated", 12, "date"),
        # --- extra tracking fields (input)
        Col("Role Category", 22, "list", "Role Category", name="Role",
            prompt="Groups job titles for the dashboard and analytics."),
        Col("First Response Date", 12, "date", name="RespDate",
            prompt="Date the company first replied (interview invite or rejection). Powers 'Avg. days to "
                   "response' and stops the NO RESPONSE flag."),
        Col("Stage When Closed", 14, "list", "Stage When Closed",
            prompt="Only for Rejected / Withdrawn: the last stage you reached (e.g. Interview 2), so interview "
                   "statistics stay correct. Leave blank if it closed right after applying."),
        # --- auto-calculated
        Col("Action Flag", 21, "auto", name="Flag", formula=(
            '=IF(OR($D{r}="",$M{r}="Rejected",$M{r}="Withdrawn",$M{r}="Accepted"),"",'
            'IF(ISNUMBER($O{r}),'
            'IF(INT($O{r})<TODAY(),"FOLLOW UP REQUIRED",IF(INT($O{r})=TODAY(),"ACTION DUE TODAY",'
            'IF(INT($O{r})-TODAY()<=Set_DueSoonDays,"DUE SOON",""))),'
            'IF(AND($M{r}="Applied",$AA{r}=""),IF(ISNUMBER($C{r}),'
            'IF(TODAY()-$C{r}>=Set_FollowUpDays,"NO RESPONSE",""),""),"")))')),
        Col("Match Category", 16, "auto", name="MatchCat", formula=(
            '=IF(ISNUMBER($L{r}),IF($L{r}>=90,"Excellent Match",IF($L{r}>=75,"Strong Match",'
            'IF($L{r}>=60,"Moderate Match","Low Match"))),"")')),
        Col("Furthest Stage", 13, "auto", name="Furthest", formula=(
            '=IF($D{r}="","",IF(OR(' + closed + '),IF($AB{r}<>"",$AB{r},'
            'IF(OR($M{r}="Rejected",ISNUMBER($C{r})),"Applied","Wishlist")),'
            'IF($M{r}="",IF(ISNUMBER($C{r}),"Applied","Wishlist"),$M{r})))')),
        Col("Days to Response", 11, "auto", name="DaysToResp", fmt="0", align="center", formula=(
            '=IF(AND(ISNUMBER($C{r}),ISNUMBER($AA{r})),$AA{r}-$C{r},"")')),
        # --- hidden helpers
        Col("Stage # (helper)", 8, "helper", name="Stage", formula=(
            '=IF($AE{r}="","",IFERROR(MATCH($AE{r},{' + stages + '},0)-1,""))')),
        Col("Action Sort Key (helper)", 8, "helper", name="ActionKey", formula=(
            '=IF($AC{r}="","",MATCH($AC{r},{' + flags + '},0)*100000'
            '+IF(ISNUMBER($O{r}),$O{r},IF(ISNUMBER($C{r}),$C{r},0))+ROW()/10000)')),
        Col("Active Sort Key (helper)", 8, "helper", name="ActiveKey", formula=(
            '=IF(AND(ISNUMBER($AG{r}),$AG{r}>=1,$AG{r}<=6,$M{r}<>"Rejected",$M{r}<>"Withdrawn"),'
            '(10-$AG{r})*100000+IF(ISNUMBER($C{r}),$C{r},0)+ROW()/10000,"")')),
    ]


def build_applications(wb, ws):
    cols = applications_columns()
    last = build_data_sheet(
        wb, ws, cols, "Applications",
        "One row per job. Status drives the pipeline · Next Action Date drives Follow-Ups · "
        "Match % is your own assessment.",
        ("D3", "White columns = your input  ·  Grey columns = auto-calculated (don't type over them)  ·  "
               "Dates as dd.mm.yyyy (Ctrl+; = today)"),
        "F5", "App",
        groups=[("Z3", "EXTRA TRACKING FIELDS  ▸  your input", ACCENT),
                ("AC3", "AUTO-CALCULATED  ▸  don't edit", MUTED)])
    put(ws, "A3", "Next ID", F(8, True, MUTED), align=Alignment(horizontal="right"))
    put(ws, "B3", "=MAX(App_ID)+1", F(10, True, ACCENT), fmt=ID_FMT, align=Alignment(horizontal="center"))
    # local copies of thresholds for conditional formatting (hidden helper columns)
    put(ws, "AG1", "Follow-up days", F(8, color=MUTED))
    put(ws, "AG2", "=Set_FollowUpDays")
    put(ws, "AH1", "Due-soon days", F(8, color=MUTED))
    put(ws, "AH2", "=Set_DueSoonDays")

    open_ = '$M{r}<>"Rejected",$M{r}<>"Withdrawn",$M{r}<>"Accepted"'.format(r=FIRST)
    status_cf(ws, f"M{FIRST}:M{last}", f"$M{FIRST}")
    for cat, lo, _, bg, fg in MATCH_BANDS:
        nxt = [b[1] for b in MATCH_BANDS if b[1] > lo]
        upper = f",$L{FIRST}<{min(nxt)}" if nxt else ""
        cf(ws, f"L{FIRST}:L{last} AD{FIRST}:AD{last}", f"AND(ISNUMBER($L{FIRST}),$L{FIRST}>={lo}{upper})",
           bg, fg, bold=cat == "Excellent Match")
    date_due_cf(ws, f"O{FIRST}:O{last}", f"$O{FIRST}", "$AH$2", open_)
    text_colour_cf(ws, f"K{FIRST}:K{last}", f"$K{FIRST}",
                   {"High": (RED_FG, True), "Medium": ("B45309", False), "Low": (MUTED, False)})
    flag_cf(ws, f"AC{FIRST}:AC{last}", f"$AC{FIRST}")
    cf(ws, f"A{FIRST}:A{last}", f"AND(ISNUMBER($A{FIRST}),COUNTIF($A${FIRST}:$A${last},$A{FIRST})>1)",
       RED_BG, RED_FG, bold=True)
    cf(ws, f"A{FIRST}:AB{last}", f'AND($D{FIRST}<>"",MOD(ROW(),2)=0)', ZEBRA, stop=False)


# ----------------------------------------------------------------------------- Companies / Contacts / Interviews
def build_companies(wb, ws):
    cols = [
        Col("Company", 24, name="Company"),
        Col("Industry", 18, "list", "Industry"),
        Col("Location", 16),
        Col("AI Relevance", 11, "list", "AI Relevance", align="center",
            prompt="How central AI is to this company: High = AI is the product / core team; Medium = growing AI "
                   "initiatives; Low = little AI work today."),
        Col("Target Role", 22, "list", "Role Category"),
        Col("Careers URL", 26, "url"),
        Col("Contact", 20),
        Col("Contact LinkedIn", 24, "url"),
        Col("Referral", 10, "list", "Yes / No", align="center"),
        Col("Priority", 10, "list", "Priority", align="center"),
        Col("Status", 15, "list", "Company Status"),
        Col("Notes", 40, "wrap"),
        Col("Applications", 12, "auto", align="center", formula=(
            '=IF($A{r}="","",COUNTIFS(App_Company,$A{r},App_Stage,">=1"))')),
        Col("Active Applications", 12, "auto", align="center", formula=(
            '=IF($A{r}="","",COUNTIFS(App_Company,$A{r},App_ActiveKey,">0"))')),
        Col("Contacts", 10, "auto", align="center", formula='=IF($A{r}="","",COUNTIF(Ct_Company,$A{r}))'),
    ]
    last = build_data_sheet(
        wb, ws, cols, "Companies",
        "Target companies — worth tracking even when there's no perfect opening yet.",
        ("A3", "Use the same company spelling as on Applications and Contacts so the counts on the right link up."),
        "B5", "Co", groups=[("M3", "AUTO-CALCULATED  ▸  from Applications & Contacts", MUTED)])
    text_colour_cf(ws, f"J{FIRST}:J{last}", f"$J{FIRST}",
                   {"High": (RED_FG, True), "Medium": ("B45309", False), "Low": (MUTED, False)})
    text_colour_cf(ws, f"D{FIRST}:D{last}", f"$D{FIRST}",
                   {"High": ("1E40AF", True), "Medium": (TEXT, False), "Low": (MUTED, False)})
    for value, style in {"Applied": "Applied", "Interviewing": "Interview 1", "Offer": "Offer",
                         "Not a Fit": "Withdrawn"}.items():
        bg, fg, _ = STATUS_STYLE[style]
        cf(ws, f"K{FIRST}:K{last}", f'$K{FIRST}="{value}"', bg, fg)
    cf(ws, f"N{FIRST}:N{last}", f"AND(ISNUMBER($N{FIRST}),$N{FIRST}>0)", fg=ACCENT, bold=True)
    cf(ws, f"A{FIRST}:L{last}", f'AND($A{FIRST}<>"",MOD(ROW(),2)=0)', ZEBRA, stop=False)


def build_contacts(wb, ws):
    cols = [
        Col("Name", 22, name="Name"),
        Col("Company", 22, name="Company"),
        Col("Job Title", 24, name="JobTitle"),
        Col("LinkedIn", 26, "url", name="LinkedIn"),
        Col("Connection Date", 12, "date"),
        Col("Relationship", 17, "list", "Relationship", name="Relationship"),
        Col("Contacted?", 11, "list", "Yes / No", name="Contacted", align="center"),
        Col("Message Sent", 12, "date", prompt="Date you sent your message. Keep the wording in Notes if useful."),
        Col("Response", 13, "list", "Response", name="Response"),
        Col("Follow-up Date", 12, "date", name="FollowUp"),
        Col("Referral Potential", 12, "list", "Referral Potential", name="RefPotential", align="center"),
        Col("Notes", 44, "wrap", name="Notes"),
        Col("Follow-up Key (helper)", 8, "helper", name="Key", formula=(
            '=IF(AND($A{r}<>"",ISNUMBER($J{r})),IF(INT($J{r})-TODAY()<=Set_DueSoonDays,'
            'INT($J{r})+ROW()/10000,""),"")')),
    ]
    last = build_data_sheet(
        wb, ws, cols, "Contacts", "Networking: recruiters, hiring managers, alumni and referrers.",
        ("A3", "Set a Follow-up Date for anyone you want to ping — due and overdue ones appear on Follow-Ups."),
        "B5", "Ct")
    put(ws, "M1", "Due-soon days", F(8, color=MUTED))
    put(ws, "M2", "=Set_DueSoonDays")
    date_due_cf(ws, f"J{FIRST}:J{last}", f"$J{FIRST}", "$M$2")
    text_colour_cf(ws, f"I{FIRST}:I{last}", f"$I{FIRST}", {
        "Positive": (GREEN_FG, True), "Meeting": (GREEN_FG, True), "Referral": (GREEN_FG, True),
        "Negative": (RED_FG, False), "No Response": (FAINT, False)})
    text_colour_cf(ws, f"K{FIRST}:K{last}", f"$K{FIRST}", {
        "High": (GREEN_FG, True), "Medium": ("B45309", False), "Low": (MUTED, False), "Unknown": (FAINT, False)})
    text_colour_cf(ws, f"G{FIRST}:G{last}", f"$G{FIRST}", {"Yes": (GREEN_FG, True), "No": (MUTED, False)})
    cf(ws, f"A{FIRST}:L{last}", f'AND($A{FIRST}<>"",MOD(ROW(),2)=0)', ZEBRA, stop=False)


def build_interviews(wb, ws):
    cols = [
        Col("Company", 22, name="Company"),
        Col("Job Title", 26, name="JobTitle"),
        Col("Interview Stage", 13, "list", "Interview Stage", name="Stage"),
        Col("Interview Date", 15, "date", name="Date", fmt=DAY_FMT),
        Col("Interviewer", 20, name="Interviewer"),
        Col("Interview Type", 16, "list", "Interview Type", name="Type"),
        Col("Preparation Status", 13, "list", "Preparation Status", name="Prep"),
        Col("Key Topics", 32, "wrap"),
        Col("Questions to Ask", 32, "wrap"),
        Col("My Performance", 14, "list", "My Performance"),
        Col("Feedback", 28, "wrap"),
        Col("Next Step", 24, "wrap", name="NextStep"),
        Col("Next Step Date", 12, "date"),
        Col("Notes", 32, "wrap"),
        Col("Upcoming Key (helper)", 8, "helper", name="Key", formula=(
            '=IF(AND($A{r}<>"",ISNUMBER($D{r})),IF(AND(INT($D{r})>=TODAY(),'
            'INT($D{r})-TODAY()<=Set_InterviewWindow),$D{r}+ROW()/100000,""),"")')),
    ]
    last = build_data_sheet(
        wb, ws, cols, "Interviews", "One row per interview round — preparation, performance and what happens next.",
        ("A3", "Interviews in the next few days are highlighted; 'Not Started' prep turns red when the interview is close."),
        "C5", "Int")
    put(ws, "O1", "Due-soon days", F(8, color=MUTED))
    put(ws, "O2", "=Set_DueSoonDays")
    d = f"$D{FIRST}"
    cf(ws, f"D{FIRST}:D{last}", f"AND(ISNUMBER({d}),INT({d})=TODAY())", "FDE68A", "78350F", bold=True)
    cf(ws, f"D{FIRST}:D{last}", f"AND(ISNUMBER({d}),INT({d})>TODAY(),INT({d})-TODAY()<=$O$2)", YELLOW_BG, YELLOW_FG)
    cf(ws, f"G{FIRST}:G{last}",
       f'AND($G{FIRST}="Not Started",ISNUMBER({d}),INT({d})>=TODAY(),INT({d})-TODAY()<=$O$2)', RED_BG, RED_FG, True)
    text_colour_cf(ws, f"G{FIRST}:G{last}", f"$G{FIRST}", {
        "Not Started": (RED_FG, False), "In Progress": ("B45309", False), "Ready": (GREEN_FG, True),
        "Completed": (FAINT, False)})
    text_colour_cf(ws, f"J{FIRST}:J{last}", f"$J{FIRST}", {
        "Poor": (RED_FG, True), "Needs Improvement": ("B45309", False), "Good": ("3F6212", False),
        "Very Good": (GREEN_FG, False), "Excellent": (GREEN_FG, True)})
    for stage in ("Screening", "Interview 1", "Interview 2", "Interview 3+", "Final"):
        bg, fg, _ = STATUS_STYLE["Interview 2" if stage == "Interview 3+" else stage]
        cf(ws, f"C{FIRST}:C{last}", f'$C{FIRST}="{stage}"', bg, fg)
    date_due_cf(ws, f"M{FIRST}:M{last}", f"$M{FIRST}", "$O$2")
    cf(ws, f"A{FIRST}:N{last}", f'AND($A{FIRST}<>"",MOD(ROW(),2)=0)', ZEBRA, stop=False)


# ----------------------------------------------------------------------------- Settings
def build_settings(wb, ws):
    ws.sheet_view.showGridLines = False
    for i in range(len(LISTS)):
        ws.column_dimensions[col_letter(LIST_FIRST_COL + i)].width = 22
    ws.column_dimensions["A"].width = 2
    put(ws, "B1", "Settings", F(16, True, INK))
    put(ws, "B2", "Change the yellow values any time — every formula, flag and highlight follows them.", F(9, color=MUTED))
    put(ws, "B4", "CONFIGURABLE VALUES", F(8, True, MUTED))
    put(ws, "E4", "VALUE", F(8, True, MUTED), align=Alignment(horizontal="center"))
    put(ws, "F4", "WHAT IT DOES", F(8, True, MUTED))
    for i, (name, label, default, lo, hi, note) in enumerate(SETTINGS):
        r = SETTINGS_ROW0 + i
        put(ws, f"B{r}", label, F(10, True), align=Alignment(vertical="center", indent=1))
        put(ws, f"E{r}", default, F(11, True, INPUT_FG), INPUT_BG, Alignment(horizontal="center", vertical="center"),
            border=Border(*(side("FDE68A"),) * 4))
        put(ws, f"F{r}", note, F(9, color=MUTED), align=Alignment(vertical="center"))
        ws.row_dimensions[r].height = 22
        add_name(wb, name, f"Settings!$E${r}")
        dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi))
        dv.errorTitle, dv.error, dv.showErrorMessage = "Whole number", f"Enter a whole number from {lo} to {hi}.", True
        ws.add_data_validation(dv)
        dv.add(f"E{r}")

    put(ws, "B10", "DROPDOWN LISTS", F(8, True, MUTED))
    put(ws, "B11", "Rename items freely — except the two lists marked FIXED, which formulas rely on. "
                   "Renaming doesn't change rows that already use the old name.", F(8, color=FAINT, italic=True))
    for i, (name, items) in enumerate(LISTS.items()):
        L = col_letter(LIST_FIRST_COL + i)
        fixed = name in FIXED_LISTS
        put(ws, f"{L}{LIST_HEAD_ROW}", name + ("  · FIXED" if fixed else ""), F(9, True, WHITE),
            "7F1D1D" if fixed else HEAD_BG, Alignment(vertical="center", indent=1))
        for j, item in enumerate(items):
            put(ws, f"{L}{LIST_FIRST_ROW + j}", item, F(10), align=Alignment(indent=1),
                border=Border(bottom=side(LINE_SOFT)))
    ws.row_dimensions[LIST_HEAD_ROW].height = 22
    note_row = LIST_FIRST_ROW + max(len(v) for v in LISTS.values()) + 1
    put(ws, f"B{note_row}", "Skill Tags isn't a dropdown: type several tags into 'Key Technologies / Skills' separated "
                            "by commas. Analytics counts the tags listed here (case-insensitive).",
        F(8, color=FAINT, italic=True))


# ----------------------------------------------------------------------------- shared report helpers
def section(ws, ref, text):
    put(ws, ref, text.upper(), F(8, True, MUTED))


def thead(ws, row, cells):
    for ref_col, text, align in cells:
        put(ws, f"{ref_col}{row}", text, F(8, True, MUTED),
            align=Alignment(horizontal=align, vertical="bottom", wrap_text=True, indent=1 if align == "left" else 0),
            border=Border(bottom=side("CBD5E1")))


def trow_border(ws, row, cols):
    for c in cols:
        ws[f"{c}{row}"].border = Border(bottom=side(LINE_SOFT))


def bar_formula(value, max_rng, width=14):
    return (f'=IF(OR(NOT(ISNUMBER({value})),MAX({max_rng})<=0),"",IF({value}<=0,"",'
            f'REPT("█",MAX(1,ROUND({value}/MAX({max_rng})*{width},0)))))')


def card(ws, col, row, label, formula, caption, accent, fmt="0"):
    c1, c2 = col, col_letter(col_idx(col) + 1)
    for r in range(row, row + 3):
        for c in (c1, c2):
            ws[f"{c}{r}"].fill = fill(CARD_BG)
            ws[f"{c}{r}"].border = Border(
                left=side(accent, "thick") if c == c1 else None, right=side(LINE) if c == c2 else None,
                top=side(LINE) if r == row else None, bottom=side(LINE) if r == row + 2 else None)
    put(ws, f"{c1}{row}", label, F(9, True, MUTED), align=Alignment(vertical="bottom", indent=1))
    put(ws, f"{c1}{row + 1}", formula, F(22, True, INK), fmt=fmt, align=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, f"{c1}{row + 2}", caption, F(8, color=FAINT), align=Alignment(vertical="top", indent=1))
    ws.row_dimensions[row].height, ws.row_dimensions[row + 1].height, ws.row_dimensions[row + 2].height = 20, 34, 18
    return f"{c1}{row + 1}"


def extraction_rows(ws, top, n, key, pos_col, fields):
    """Rows top..top+n-1 list the records whose `key` is a number, smallest key first.

    fields: (column, kind, named range) with kind text | date | link | overdue | since | until | combo
    """
    for k in range(1, n + 1):
        r = top + k - 1
        p = f"${pos_col}{r}"
        put(ws, f"{pos_col}{r}", f'=IFERROR(MATCH(SMALL({key},{k}),{key},0),"")')
        for col, kind, rng in fields:
            v = f"INDEX({rng},{p})"
            if kind == "text":
                f = f'=IF({p}="","",{v}&"")'
            elif kind == "date":
                f = f'=IF({p}="","",IF(ISNUMBER({v}),{v},""))'
            elif kind == "link":
                f = f'=IF({p}="","",IF({v}="","",HYPERLINK({v},"Open ↗")))'
            elif kind == "overdue":
                f = f'=IF({p}="","",IF(ISNUMBER({v}),IF(INT({v})<TODAY(),TODAY()-INT({v}),""),""))'
            elif kind == "since":
                f = f'=IF({p}="","",IF(ISNUMBER({v}),TODAY()-INT({v}),""))'
            elif kind == "until":
                f = f'=IF({p}="","",IF(ISNUMBER({v}),INT({v})-TODAY(),""))'
            elif kind == "combo":  # company — job title
                f = f'=IF({p}="","",INDEX(App_Company,{p})&IF(INDEX(App_JobTitle,{p})="",""," — "&INDEX(App_JobTitle,{p})))'
            else:
                raise ValueError(kind)
            c = put(ws, f"{col}{r}", f, F(10, color=ACCENT if kind == "link" else TEXT))
            if kind == "date":
                c.number_format = DATE_FMT
            c.alignment = Alignment(horizontal="center" if kind in ("date", "overdue", "since", "until") else "left",
                                    indent=0 if kind in ("date", "overdue", "since", "until") else 1)
        ws.row_dimensions[r].height = 18
    ws.column_dimensions[pos_col].hidden = True
    # row rules only on filled rows, so an empty list reads as white space rather than a blank grid
    cols = sorted(col_idx(f[0]) for f in fields)
    ws.conditional_formatting.add(
        f"{col_letter(cols[0])}{top}:{col_letter(cols[-1])}{top + n - 1}",
        FormulaRule(formula=[f'${pos_col}{top}<>""'], border=Border(bottom=side(LINE)), stopIfTrue=False))


def more_line(ws, ref, key, n, where):
    put(ws, ref, f'=IF(COUNT({key})>{n},"+ "&(COUNT({key})-{n})&" more — {where}","")', F(8, color=MUTED, italic=True))


# ----------------------------------------------------------------------------- Follow-Ups
def build_followups(wb, ws, refs):
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 22, "C": 24, "D": 30, "E": 14, "F": 28, "G": 13, "H": 12, "I": 15, "J": 11, "K": 10,
              "L": 22, "M": 9, "N": 6}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    put(ws, "B1", "Follow-Ups", F(18, True, INK))
    put(ws, "B2", "Everything that needs your attention, most urgent first. Fills itself from Applications, "
                  "Contacts and Interviews.", F(9, color=MUTED))
    put(ws, "B3", '="No-response threshold: "&Set_FollowUpDays&" days  ·  Due-soon window: "&Set_DueSoonDays&'
                  '" days  ·  Upcoming interviews: next "&Set_InterviewWindow&" days  ·  change these on Settings"',
        F(8, color=FAINT, italic=True))
    ws.row_dimensions[1].height = 30

    kpis = [("B", "Follow Up Required", '=COUNTIF(App_Flag,"FOLLOW UP REQUIRED")', "overdue"),
            ("C", "Action Due Today", '=COUNTIF(App_Flag,"ACTION DUE TODAY")', "today"),
            ("D", "No Response Yet", '=COUNTIF(App_Flag,"NO RESPONSE")', "noresp"),
            ("E", "Due Soon", '=COUNTIF(App_Flag,"DUE SOON")', "soon"),
            ("F", "Contact Follow-ups", "=COUNT(Ct_Key)", "contacts"),
            ("G", "Upcoming Interviews", "=COUNT(Int_Key)", "interviews")]
    for col, label, formula, key in kpis:
        put(ws, f"{col}5", label, F(8, True, MUTED), CARD_BG, Alignment(vertical="bottom", wrap_text=True, indent=1),
            border=Border(top=side(LINE)))
        put(ws, f"{col}6", formula, F(20, True, INK), CARD_BG, Alignment(vertical="center", horizontal="left", indent=1),
            border=Border(bottom=side(LINE)))
        refs[f"fu.{key}"] = ("Follow-Ups", f"{col}6")
    ws.row_dimensions[5].height, ws.row_dimensions[6].height = 26, 34
    cf(ws, "B6", "$B$6>0", fg="B91C1C")
    cf(ws, "C6", "$C$6>0", fg="B45309")

    # applications action list
    n_app, top = 25, 10
    section(ws, "B8", "Applications — action list (overdue first, then due today, no response, due soon)")
    thead(ws, 9, [("B", "Action Flag", "left"), ("C", "Company", "left"), ("D", "Job Title", "left"),
                  ("E", "Status", "left"), ("F", "Next Action", "left"), ("G", "Next Action Date", "center"),
                  ("H", "Days Overdue", "center"), ("I", "Date Applied", "center"), ("J", "Days Since Applied", "center"),
                  ("K", "Priority", "left"), ("L", "Recruiter / Hiring Manager", "left"), ("M", "Job Link", "left")])
    ws.row_dimensions[9].height = 28
    extraction_rows(ws, top, n_app, "App_ActionKey", "N", [
        ("B", "text", "App_Flag"), ("C", "text", "App_Company"), ("D", "text", "App_JobTitle"),
        ("E", "text", "App_Status"), ("F", "text", "App_NextAction"), ("G", "date", "App_NextDate"),
        ("H", "overdue", "App_NextDate"), ("I", "date", "App_DateApplied"), ("J", "since", "App_DateApplied"),
        ("K", "text", "App_Priority"), ("L", "text", "App_Contact"), ("M", "link", "App_URL")])
    end = top + n_app - 1
    flag_cf(ws, f"B{top}:B{end}", f"$B{top}")
    status_cf(ws, f"E{top}:E{end}", f"$E{top}")
    cf(ws, f"H{top}:H{end}", f"AND(ISNUMBER($H{top}),$H{top}>0)", fg="B91C1C", bold=True)
    text_colour_cf(ws, f"K{top}:K{end}", f"$K{top}", {"High": (RED_FG, True), "Low": (MUTED, False)})
    more_line(ws, f"B{end + 1}", "App_ActionKey", n_app, "filter the Action Flag column on Applications")
    refs["fu.app_list"] = ("Follow-Ups", top, n_app)

    # interviews
    n_iv, top = 10, end + 5
    put(ws, f"B{top - 2}", '="UPCOMING INTERVIEWS — NEXT "&Set_InterviewWindow&" DAYS"', F(8, True, MUTED))
    thead(ws, top - 1, [("B", "Interview Date", "left"), ("C", "Company", "left"), ("D", "Job Title", "left"),
                        ("E", "Stage", "left"), ("F", "Interview Type", "left"), ("G", "Days Until", "center"),
                        ("H", "Preparation", "left"), ("I", "Interviewer", "left"), ("J", "", "left"),
                        ("K", "", "left"), ("L", "Next Step", "left")])
    ws.row_dimensions[top - 1].height = 28
    extraction_rows(ws, top, n_iv, "Int_Key", "N", [
        ("B", "date", "Int_Date"), ("C", "text", "Int_Company"), ("D", "text", "Int_JobTitle"),
        ("E", "text", "Int_Stage"), ("F", "text", "Int_Type"), ("G", "until", "Int_Date"),
        ("H", "text", "Int_Prep"), ("I", "text", "Int_Interviewer"), ("L", "text", "Int_NextStep")])
    end = top + n_iv - 1
    for r in range(top, end + 1):
        ws[f"B{r}"].number_format = DAY_FMT
        ws[f"B{r}"].alignment = Alignment(horizontal="left", indent=1)
    cf(ws, f"G{top}:G{end}", f"AND(ISNUMBER($G{top}),$G{top}=0)", "FDE68A", "78350F", bold=True)
    text_colour_cf(ws, f"H{top}:H{end}", f"$H{top}", {"Not Started": (RED_FG, True), "In Progress": ("B45309", False),
                                                      "Ready": (GREEN_FG, True)})
    more_line(ws, f"B{end + 1}", "Int_Key", n_iv, "see the Interviews tab")
    refs["fu.iv_list"] = ("Follow-Ups", top, n_iv)

    # contacts
    n_ct, top = 15, end + 5
    section(ws, f"B{top - 2}", "Contacts — follow-ups due (overdue, today and within the due-soon window)")
    thead(ws, top - 1, [("B", "Name", "left"), ("C", "Company", "left"), ("D", "Job Title", "left"),
                        ("E", "Response", "left"), ("F", "Notes", "left"), ("G", "Follow-up Date", "center"),
                        ("H", "Days Overdue", "center"), ("I", "Referral Potential", "left"), ("J", "Contacted?", "left"),
                        ("K", "", "left"), ("L", "Relationship", "left"), ("M", "LinkedIn", "left")])
    ws.row_dimensions[top - 1].height = 28
    extraction_rows(ws, top, n_ct, "Ct_Key", "N", [
        ("B", "text", "Ct_Name"), ("C", "text", "Ct_Company"), ("D", "text", "Ct_JobTitle"),
        ("E", "text", "Ct_Response"), ("F", "text", "Ct_Notes"), ("G", "date", "Ct_FollowUp"),
        ("H", "overdue", "Ct_FollowUp"), ("I", "text", "Ct_RefPotential"), ("J", "text", "Ct_Contacted"),
        ("L", "text", "Ct_Relationship"), ("M", "link", "Ct_LinkedIn")])
    end = top + n_ct - 1
    cf(ws, f"G{top}:G{end}", f"AND(ISNUMBER($G{top}),$G{top}<TODAY())", RED_BG, RED_FG, bold=True)
    cf(ws, f"G{top}:G{end}", f"ISNUMBER($G{top})", YELLOW_BG, YELLOW_FG)
    more_line(ws, f"B{end + 1}", "Ct_Key", n_ct, "sort the Contacts tab by Follow-up Date")
    refs["fu.ct_list"] = ("Follow-Ups", top, n_ct)

    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ----------------------------------------------------------------------------- Analytics
def build_analytics(wb, ws, refs):
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 44, "C": 14, "D": 14, "E": 14, "F": 12, "G": 20, "H": 8}.items():
        ws.column_dimensions[c].width = w
    ws.column_dimensions["H"].hidden = True
    put(ws, "B1", "Analytics", F(18, True, INK))
    put(ws, "B2", "What's working? Interview = reached Interview 1 or later (including applications later rejected). "
                  "Rates use submitted applications (status beyond Wishlist).", F(9, color=MUTED))
    put(ws, "B3", '="A group is only called \'best\' once it has at least "&Set_MinSample&'
                  '" applications (Settings) — small samples are noise."', F(8, color=FAINT, italic=True))
    ws.row_dimensions[1].height = 30

    # summary
    section(ws, "B5", "Summary")
    summary = [
        ("Applications submitted", '=COUNTIFS(App_Stage,">=1")', "0", "an.total"),
        ("Reached interview (Interview 1+)", '=COUNTIFS(App_Stage,">=3")', "0", "an.ints"),
        ("Application → Interview rate", '=IF(C6=0,"–",C7/C6)', PCT_FMT, "an.rate"),
        ("Avg. Match % · applications that reached interview",
         '=IFERROR(AVERAGEIFS(App_Match,App_Stage,">=3"),"–")', '0"%"', "an.match_int"),
        ("Avg. Match % · applications without an interview",
         '=IFERROR(AVERAGEIFS(App_Match,App_Stage,">=1",App_Stage,"<=2"),"–")', '0"%"', "an.match_noint"),
    ]
    for i, (label, formula, fmt, key) in enumerate(summary):
        r = 6 + i
        put(ws, f"B{r}", label, F(10), align=Alignment(indent=1))
        put(ws, f"C{r}", formula, F(11, True, INK), fmt=fmt, align=Alignment(horizontal="center"))
        trow_border(ws, r, "BC")
        refs[key] = ("Analytics", f"C{r}")

    row = 13

    def group_table(row, title, first_header, labels, crit_rng, key, callout=True, bar="rate"):
        """labels: list of (display formula/text, criteria expression or None for 'Not set')."""
        section(ws, f"B{row}", title)
        h = row + 1
        thead(ws, h, [("B", first_header, "left"), ("C", "Applications", "center"), ("D", "Interviews", "center"),
                      ("E", "Interview Rate", "center"), ("F", "Offers", "center"),
                      ("G", "Interview rate" if bar == "rate" else "Frequency", "left")])
        ws.row_dimensions[h].height = 26
        r0 = h + 1
        rows = []
        for i, (label, crit) in enumerate(labels):
            r = r0 + i
            rows.append(r)
            put(ws, f"B{r}", label, F(10, color=MUTED if crit is None else TEXT), align=Alignment(indent=1))
            if crit is None:  # remainder row: everything not covered by the rows above
                a, b = r0, r - 1
                put(ws, f"C{r}", f"=MAX(0,$C$6-SUM(C{a}:C{b}))")
                put(ws, f"D{r}", f"=MAX(0,$C$7-SUM(D{a}:D{b}))")
                put(ws, f"F{r}", f'=MAX(0,COUNTIFS(App_Stage,">=6")-SUM(F{a}:F{b}))')
                put(ws, f"H{r}", -1)
            elif crit_rng == "skills":
                tag = (f'ISNUMBER(SEARCH(","&SUBSTITUTE($B{r}," ","")&",",'
                       f'","&SUBSTITUTE(SUBSTITUTE(App_Skills," ",""),";",",")&","))')
                put(ws, f"C{r}", f"=SUMPRODUCT({tag}*ISNUMBER(App_Stage)*(App_Stage>=1))")
                put(ws, f"D{r}", f"=SUMPRODUCT({tag}*ISNUMBER(App_Stage)*(App_Stage>=3))")
                put(ws, f"F{r}", f"=SUMPRODUCT({tag}*ISNUMBER(App_Stage)*(App_Stage>=6))")
                put(ws, f"H{r}", -1)
            else:
                put(ws, f"C{r}", f'=COUNTIFS({crit_rng},{crit},App_Stage,">=1")')
                put(ws, f"D{r}", f'=COUNTIFS({crit_rng},{crit},App_Stage,">=3")')
                put(ws, f"F{r}", f'=COUNTIFS({crit_rng},{crit},App_Stage,">=6")')
                put(ws, f"H{r}", f"=IF(AND(C{r}>0,C{r}>=Set_MinSample),D{r}/C{r},-1)")
            put(ws, f"E{r}", f'=IF(C{r}=0,"–",D{r}/C{r})', fmt=PCT_FMT)
            for c in "CDEF":
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
                ws[f"{c}{r}"].font = F(10)
            src = f"E{r}" if bar == "rate" else f"C{r}"
            rng = f"$E${r0}:$E${r0 + len(labels) - 1}" if bar == "rate" else f"$C${r0}:$C${r0 + len(labels) - 1}"
            if bar == "rate":
                put(ws, f"G{r}", f'=IF(ISNUMBER({src}),IF({src}>0,REPT("█",MAX(1,ROUND({src}*14,0))),""),"")',
                    F(9, color=ACCENT_SOFT))
            else:
                put(ws, f"G{r}", bar_formula(src, rng), F(9, color=ACCENT_SOFT))
            trow_border(ws, r, "BCDEFG")
        last = rows[-1]
        refs[key] = ("Analytics", r0, len(labels))
        nxt = last + 1
        if callout:
            hr = f"$H${r0}:$H${last}"
            best = f"MATCH(MAX({hr}),{hr},0)"
            put(ws, f"B{nxt}", "Best performer", F(9, True, MUTED), align=Alignment(indent=1))
            put(ws, f"C{nxt}",
                f'=IF(MAX({hr})<0,"Not enough data yet — min. "&Set_MinSample&" applications per group",'
                f'IF(MAX({hr})=0,"No interviews yet",INDEX($B${r0}:$B${last},{best})&"  —  "&ROUND(MAX({hr})*100,0)'
                f'&"% interview rate ("&INDEX($D${r0}:$D${last},{best})&" of "&INDEX($C${r0}:$C${last},{best})&")"))',
                F(10, True, GREEN_FG))
            refs[key + ".best"] = ("Analytics", f"C{nxt}")
            nxt += 1
        return nxt + 2

    def labels_from(list_name, first_row):
        """Row labels pulled from a Settings list; each row's criteria is its own label cell."""
        n = len(LISTS[list_name])
        return [(f"={list_cell(list_name, i)}", f"$B{first_row + i}") for i in range(n)] + [("Not set", None)]

    for title, header, list_name, rng, key in [
            ("Interview rate by CV version", "CV Version", "CV Version", "App_CV", "an.cv"),
            ("Interview rate by role category", "Role Category", "Role Category", "App_Role", "an.role"),
            ("Interview rate by source", "Source", "Source", "App_Source", "an.source")]:
        row = group_table(row, title, header, labels_from(list_name, row + 2), rng, key)
    match_labels = [(f"{cat} ({lab})", f'"{cat}"') for cat, _, lab, _, _ in MATCH_BANDS] + [("Not rated", None)]
    match_start = row + 2
    row = group_table(row, "Interview rate by Match % (your assessment)", "Match Category", match_labels,
                      "App_MatchCat", "an.match")
    for i, (cat, _, _, bg, fg) in enumerate(MATCH_BANDS):
        ws[f"B{match_start + i}"].fill = fill(bg)
        ws[f"B{match_start + i}"].font = F(10, True, fg)

    # weekly activity
    section(ws, f"B{row}", "Weekly activity — last 12 weeks (Monday to Sunday)")
    h = row + 1
    thead(ws, h, [("B", "Calendar week", "left"), ("C", "Week of (Mon)", "center"), ("D", "Applications", "center"),
                  ("E", "Interviews held", "center"), ("F", "", "center"), ("G", "Applications", "left")])
    ws.row_dimensions[h].height = 26
    w0 = h + 1
    for k in range(12):
        r = w0 + k
        put(ws, f"C{r}", f"=TODAY()-WEEKDAY(TODAY(),2)+1-7*{11 - k}", F(10), fmt=DATE_FMT,
            align=Alignment(horizontal="center"))
        put(ws, f"B{r}", f'="KW "&WEEKNUM(C{r},21)&IF(C{r}=TODAY()-WEEKDAY(TODAY(),2)+1,"  (this week)","")',
            F(10), align=Alignment(indent=1))
        put(ws, f"D{r}", f'=COUNTIFS(App_DateApplied,">="&C{r},App_DateApplied,"<"&(C{r}+7))', F(10),
            align=Alignment(horizontal="center"))
        put(ws, f"E{r}", f'=COUNTIFS(Int_Date,">="&C{r},Int_Date,"<"&(C{r}+7))', F(10),
            align=Alignment(horizontal="center"))
        put(ws, f"G{r}", bar_formula(f"D{r}", f"$D${w0}:$D${w0 + 11}"), F(9, color=ACCENT_SOFT))
        trow_border(ws, r, "BCDEFG")
    put(ws, f"B{w0 + 12}", "Average per week", F(9, True, MUTED), align=Alignment(indent=1))
    put(ws, f"D{w0 + 12}", f"=AVERAGE(D{w0}:D{w0 + 11})", F(10, True), fmt="0.0", align=Alignment(horizontal="center"))
    put(ws, f"E{w0 + 12}", f"=AVERAGE(E{w0}:E{w0 + 11})", F(10, True), fmt="0.0", align=Alignment(horizontal="center"))
    refs["an.weekly"] = ("Analytics", w0, 12)
    row = w0 + 15

    # skills
    skills = [(f"={list_cell('Skill Tags', i)}", "skill") for i in range(len(LISTS["Skill Tags"]))]
    row = group_table(row, "Skill tags in your applications (from 'Key Technologies / Skills')", "Skill Tag",
                      skills, "skills", "an.skills", callout=False, bar="count")

    # active applications
    n_act = 30
    section(ws, f"B{row}", "Active applications — furthest along first")
    h = row + 1
    thead(ws, h, [("B", "Company — Job Title", "left"), ("C", "Status", "left"), ("D", "Date Applied", "center"),
                  ("E", "Days Since Applied", "center"), ("F", "Priority", "left"), ("G", "Next Action Date", "center")])
    ws.row_dimensions[h].height = 26
    top = h + 1
    extraction_rows(ws, top, n_act, "App_ActiveKey", "H", [
        ("B", "combo", None), ("C", "text", "App_Status"), ("D", "date", "App_DateApplied"),
        ("E", "since", "App_DateApplied"), ("F", "text", "App_Priority"), ("G", "date", "App_NextDate")])
    status_cf(ws, f"C{top}:C{top + n_act - 1}", f"$C{top}")
    more_line(ws, f"B{top + n_act}", "App_ActiveKey", n_act, "filter Status on Applications")
    refs["an.active"] = ("Analytics", top, n_act)


# ----------------------------------------------------------------------------- Dashboard
def build_dashboard(wb, ws, refs):
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 16, "C": 15, "D": 11, "E": 16, "F": 3, "G": 16, "H": 15, "I": 11, "J": 16, "K": 2,
                 "L": 6, "M": 6}.items():
        ws.column_dimensions[c].width = w
    ws.column_dimensions["L"].hidden = True
    ws.column_dimensions["M"].hidden = True

    put(ws, "B1", "Job Search Dashboard", F(20, True, INK), align=Alignment(vertical="center"))
    put(ws, "I1", "As of", F(8, True, MUTED), align=Alignment(horizontal="right", vertical="center"))
    put(ws, "J1", "=TODAY()", F(10, True, TEXT), fmt=DAY_FMT, align=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B2", "AI Engineer · Generative AI · Agentic AI · AI Automation · AI Solutions  —  Stuttgart region, "
                  "Germany & remote", F(9, color=MUTED))
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[3].height = 8

    section(ws, "B4", "Overview")
    total = card(ws, "B", 5, "Total Applications", '=COUNTIFS(App_Stage,">=1")', "submitted (beyond Wishlist)", ACCENT)
    card(ws, "D", 5, "Applications This Week",
         '=COUNTIFS(App_DateApplied,">="&(TODAY()-WEEKDAY(TODAY(),2)+1),App_DateApplied,"<"&(TODAY()-WEEKDAY(TODAY(),2)+8))',
         "Mon–Sun, by Date Applied", ACCENT)
    card(ws, "G", 5, "Applications This Month",
         '=COUNTIFS(App_DateApplied,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),App_DateApplied,"<"&(EOMONTH(TODAY(),0)+1))',
         "calendar month, by Date Applied", ACCENT)
    card(ws, "I", 5, "Active Applications",
         '=COUNTIFS(App_Stage,">=1",App_Stage,"<=6",App_Status,"<>Rejected",App_Status,"<>Withdrawn")',
         "Applied → Offer, still open", ACCENT)
    ws.row_dimensions[8].height = 8
    ints = card(ws, "B", 9, "Interviews", '=COUNTIFS(App_Stage,">=3")', "reached Interview 1 or later",
                STATUS_STYLE["Interview 1"][2])
    finals = card(ws, "D", 9, "Final Rounds", '=COUNTIFS(App_Stage,">=5")', "reached the final round",
                  STATUS_STYLE["Final"][2])
    offers = card(ws, "G", 9, "Offers", '=COUNTIFS(App_Stage,">=6")', "offers received", STATUS_STYLE["Offer"][2])
    card(ws, "I", 9, "Accepted", '=COUNTIF(App_Status,"Accepted")', "offers accepted", STATUS_STYLE["Accepted"][2])
    for key, ref in zip(["total", "week", "month", "active", "ints", "finals", "offers", "accepted"],
                        ["B6", "D6", "G6", "I6", "B10", "D10", "G10", "I10"]):
        refs[f"dash.{key}"] = ("Job Search Dashboard", ref)

    # today banner
    ws.row_dimensions[12].height = 10
    fu = "'Follow-Ups'!"
    put(ws, "L13", f"={fu}$B$6")
    put(ws, "M13", f"=SUM({fu}$B$6:$G$6)")
    put(ws, "B13",
        f'=IF($M$13=0,"✓  Nothing due — all follow-ups are on track.","Today:  "&{fu}$B$6&" overdue · "&{fu}$C$6'
        f'&" due today · "&{fu}$D$6&" no response · "&{fu}$E$6&" due soon · "'
        f'&{fu}$F$6&" contact follow-ups · "&{fu}$G$6&" upcoming interviews   →  Follow-Ups tab")',
        F(10, True, TEXT), align=Alignment(vertical="center", indent=1))
    for c in "BCDEFGHIJ":
        ws[f"{c}13"].fill = fill(CARD_BG)
    cf(ws, "B13:J13", "$L$13>0", "FEF2F2", "991B1B", bold=True)
    cf(ws, "B13:J13", "$M$13>0", "FFFBEB", "92400E", bold=True)
    cf(ws, "B13:J13", "$M$13=0", "F0FDF4", GREEN_FG, bold=True)
    ws.row_dimensions[13].height = 28
    ws.row_dimensions[14].height = 12

    # pipeline
    section(ws, "B15", "Pipeline · current status")
    thead(ws, 16, [("B", "Status", "left"), ("C", "", "left"), ("D", "Count", "center"), ("E", "", "left")])
    p0 = 17
    statuses = LISTS["Status"]
    for i, s in enumerate(statuses):
        r = p0 + i
        bg, fg, bar = STATUS_STYLE[s]
        put(ws, f"B{r}", s, F(9, True, fg), bg, Alignment(horizontal="left", vertical="center", indent=1))
        put(ws, f"D{r}", f"=COUNTIF(App_Status,$B{r})", F(10, True), align=Alignment(horizontal="center"))
        put(ws, f"E{r}", bar_formula(f"D{r}", f"$D${p0}:$D${p0 + len(statuses) - 1}"), F(9, color=bar))
        trow_border(ws, r, "CDE")
        ws.row_dimensions[r].height = 18
    refs["dash.pipeline"] = ("Job Search Dashboard", p0, len(statuses))

    # conversion
    section(ws, "G15", "Conversion")
    thead(ws, 16, [("G", "Metric", "left"), ("H", "", "left"), ("I", "Rate", "center"), ("J", "Basis", "left")])
    conv = [
        ("Application → Interview", f'=IF({total}=0,"–",{ints}/{total})', f'={ints}&" of "&{total}', "app_int"),
        ("Interview → Final", f'=IF({ints}=0,"–",{finals}/{ints})', f'={finals}&" of "&{ints}', "int_final"),
        ("Final → Offer", f'=IF({finals}=0,"–",{offers}/{finals})', f'={offers}&" of "&{finals}', "final_offer"),
        ("Overall Offer Rate", f'=IF({total}=0,"–",{offers}/{total})', f'={offers}&" of "&{total}', "offer_rate"),
        ("Avg. Days to First Response", '=IF(COUNT(App_DaysToResp)=0,"–",AVERAGE(App_DaysToResp))',
         '="from "&COUNT(App_DaysToResp)&" replies"', "avg_days"),
    ]
    for i, (label, value, basis, key) in enumerate(conv):
        r = p0 + i
        put(ws, f"G{r}", label, F(10), align=Alignment(indent=1))
        put(ws, f"I{r}", value, F(11, True, INK), fmt="0.0" if key == "avg_days" else "0.0%",
            align=Alignment(horizontal="center"))
        put(ws, f"J{r}", basis, F(8, color=FAINT), align=Alignment(indent=1))
        trow_border(ws, r, "GHIJ")
        refs[f"dash.{key}"] = ("Job Search Dashboard", f"I{r}")

    section(ws, "G23", "Data checks")
    checks = [
        ("Submitted, but no Date Applied",
         "=SUMPRODUCT(ISNUMBER(App_Stage)*(App_Stage>=1)*NOT(ISNUMBER(App_DateApplied)))", "missing_date"),
        ("Submitted, but no Role Category", '=SUMPRODUCT(ISNUMBER(App_Stage)*(App_Stage>=1)*(App_Role=""))',
         "missing_role"),
        ("Duplicate Application IDs", "=SUMPRODUCT(ISNUMBER(App_ID)*(COUNTIF(App_ID,App_ID)>1))", "dup_ids"),
    ]
    for i, (label, formula, key) in enumerate(checks):
        r = 24 + i
        put(ws, f"G{r}", label, F(9, color=MUTED), align=Alignment(indent=1))
        put(ws, f"I{r}", formula, F(10, True), align=Alignment(horizontal="center"))
        put(ws, f"J{r}", f'=IF(I{r}=0,"✓ ok","⚠ fix on Applications")', F(8, color=FAINT), align=Alignment(indent=1))
        trow_border(ws, r, "GHIJ")
        refs[f"dash.{key}"] = ("Job Search Dashboard", f"I{r}")
    cf(ws, "I24:J26", "$I24>0", fg="B45309", bold=True)

    # role / source
    def breakdown(col_label, col_count, col_bar, row, title, items, crit_rng, key, colors=None, crit_values=None):
        section(ws, f"{col_label}{row}", title)
        thead(ws, row + 1, [(col_label, "", "left"), (col_letter(col_idx(col_label) + 1), "", "left"),
                            (col_count, "Apps", "center"), (col_bar, "", "left")])
        r0 = row + 2
        for i, item in enumerate(items + ["Not set"]):
            r = r0 + i
            is_rest = i == len(items)
            label_font = F(10, color=MUTED if is_rest else TEXT)
            put(ws, f"{col_label}{r}", item, label_font, align=Alignment(indent=1))
            if colors and not is_rest:
                bg, fg = colors[i]
                ws[f"{col_label}{r}"].fill = fill(bg)
                ws[f"{col_label}{r}"].font = F(9, True, fg)
            if is_rest:
                f = f"=MAX(0,{total}-SUM({col_count}{r0}:{col_count}{r - 1}))"
            else:
                crit = crit_values[i] if crit_values else f"${col_label}{r}"
                f = f'=COUNTIFS({crit_rng},{crit},App_Stage,">=1")'
            put(ws, f"{col_count}{r}", f, F(10, True), align=Alignment(horizontal="center"))
            last = r0 + len(items)
            put(ws, f"{col_bar}{r}", bar_formula(f"{col_count}{r}", f"${col_count}${r0}:${col_count}${last}"),
                F(9, color=FAINT if is_rest else ACCENT_SOFT))
            trow_border(ws, r, [col_letter(col_idx(col_label) + 1), col_count, col_bar] + ([] if colors else [col_label]))
            ws.row_dimensions[r].height = 18
        refs[key] = ("Job Search Dashboard", r0, len(items) + 1, col_count)
        return r0 + len(items) + 1

    roles = [f"={list_cell('Role Category', i)}" for i in range(len(LISTS["Role Category"]))]
    sources = [f"={list_cell('Source', i)}" for i in range(len(LISTS["Source"]))]
    end_l = breakdown("B", "D", "E", 28, "Applications by role", roles, "App_Role", "dash.role")
    end_r = breakdown("G", "I", "J", 28, "Applications by source", sources, "App_Source", "dash.source")
    row = max(end_l, end_r) + 1

    match_items = [f"{cat} · {lab}" for cat, _, lab, _, _ in MATCH_BANDS]
    end_l = breakdown("B", "D", "E", row, "Match quality · your own assessment", match_items, "App_MatchCat",
                      "dash.match", colors=[(b[3], b[4]) for b in MATCH_BANDS],
                      crit_values=[f'"{b[0]}"' for b in MATCH_BANDS])
    ws[f"B{end_l - 1}"].value = "Not rated"

    section(ws, f"G{row}", "Top performers · interview rate")
    thead(ws, row + 1, [("G", "", "left"), ("H", "", "left"), ("I", "", "left"), ("J", "", "left")])
    for i, (label, key) in enumerate([("CV version", "an.cv"), ("Source", "an.source"), ("Role category", "an.role"),
                                      ("Match band", "an.match")]):
        r = row + 2 + i
        put(ws, f"G{r}", label, F(9, True, MUTED), align=Alignment(indent=1))
        put(ws, f"H{r}", f"=Analytics!{refs[key + '.best'][1]}", F(9, True, GREEN_FG))
        trow_border(ws, r, "GHIJ")
        ws.row_dimensions[r].height = 18
        refs[f"dash.best.{key}"] = ("Job Search Dashboard", f"H{r}")
    put(ws, f"G{row + 6}", '="Groups need ≥ "&Set_MinSample&" applications · details on Analytics"',
        F(8, color=FAINT, italic=True), align=Alignment(indent=1))

    # weekly chart
    row = end_l + 1
    section(ws, f"B{row}", "Applications per week · last 12 weeks")
    _, w0, n = refs["an.weekly"]
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = "col", "clustered", 55
    ch.legend = None
    ch.add_data(Reference(wb["Analytics"], min_col=4, min_row=w0, max_row=w0 + n - 1), titles_from_data=False)
    ch.set_categories(Reference(wb["Analytics"], min_col=2, min_row=w0, max_row=w0 + n - 1))
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=ACCENT, ln=LineProperties(solidFill=ACCENT))
    s.dLbls = DataLabelList(numFmt="0;-0;;", showVal=True)  # zero weeks get no label
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = s.dLbls.showPercent = False
    s.dLbls.position = "outEnd"
    s.dLbls.txPr = chart_text(900, TEXT, True)
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    ch.y_axis.scaling.min = 0
    ch.x_axis.delete = False
    ch.x_axis.txPr = chart_text(800, MUTED)
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="CBD5E1"))
    ch.x_axis.majorTickMark = "none"
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.width, ch.height = 23.5, 6.5
    ws.add_chart(ch, f"B{row + 1}")

    put(ws, f"B{row + 15}", "New here? The Guide tab explains every column, colour and metric.",
        F(8, color=FAINT, italic=True))
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ----------------------------------------------------------------------------- Guide
EXAMPLE = [
    ("Application ID", "1  → shows as APP-001", "Type the 'Next ID' shown above the Applications header. It moves with the row when you sort."),
    ("Date Added", dt.date(2026, 10, 1), "When you saved the job. Ctrl+; inserts today's date."),
    ("Date Applied", dt.date(2026, 10, 2), "Leave blank while the job is on your Wishlist. Drives weekly counts and NO RESPONSE alerts."),
    ("Company", "Example GmbH (sample)", "Use the same spelling on Companies and Contacts so their counts link up."),
    ("Job Title", "Generative AI Engineer (m/w/d)", "Exact title from the posting."),
    ("Job Type", "Full-time", "Dropdown."),
    ("Location", "Stuttgart", "City or region."),
    ("Remote / Hybrid / On-site", "Hybrid", "Dropdown."),
    ("Job URL", "https://example.com/careers/genai-engineer", "Paste the posting link — it's clickable from Follow-Ups."),
    ("Source", "LinkedIn", "Where you found it. Analytics compares interview rates per source."),
    ("Priority", "High", "High / Medium / Low."),
    ("Match %", "85  → shows as 85%, Strong Match", "Your own fit estimate, 0–100. Not an objective score."),
    ("Status", "Applied", "Current pipeline stage (see the status legend above)."),
    ("Next Action", "Follow up with the recruiter if no reply", "The single next thing you'll do."),
    ("Next Action Date", dt.date(2026, 10, 9), "Red when overdue, yellow when due soon. Move or clear it once done."),
    ("Recruiter / Hiring Manager", "Name of the recruiter", ""),
    ("Recruiter LinkedIn", "https://www.linkedin.com/in/…", ""),
    ("CV Version", "Generative AI CV", "Analytics shows which CV version gets the best interview rate."),
    ("Cover Letter", "Yes", "Yes / No / Not Required."),
    ("Referral", "No", "Yes / No."),
    ("Salary / Salary Range", "€60–70k (estimate)", "Free text."),
    ("Key Technologies / Skills", "LLMs, RAG, Python, Vector Databases, AI Agents",
     "Several tags separated by commas. Tags from Settings → Skill Tags are counted in Analytics."),
    ("Why Me?", "Master's thesis focused on Agentic AI and intelligent knowledge management, with software background "
                "and experience with LLM/RAG concepts.",
     "1–2 sentences. Reuse for cover letters, recruiter messages and 'why are you a good fit?' answers."),
    ("Notes", "Team builds internal GenAI assistants; mention thesis on knowledge management.", ""),
    ("Last Updated", dt.date(2026, 10, 2), "Update whenever anything changes (Ctrl+;)."),
    ("Role Category", "Generative AI Engineer", "Extra field: groups job titles for the dashboard and analytics."),
    ("First Response Date", "(blank until they reply)", "Extra field: date of the first reply — invite or rejection."),
    ("Stage When Closed", "(blank — only for Rejected / Withdrawn)",
     "Extra field: e.g. Interview 2 if you were rejected after the second interview. Keeps interview stats honest."),
    ("Action Flag", "(auto)", "FOLLOW UP REQUIRED / ACTION DUE TODAY / NO RESPONSE / DUE SOON."),
    ("Match Category", "(auto) Strong Match", "From Match %."),
    ("Furthest Stage", "(auto) Applied", "Status, or Stage When Closed for closed applications."),
    ("Days to Response", "(auto)", "First Response Date − Date Applied."),
]


def build_guide(ws):
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 2, "B": 28, "C": 58, "D": 70}.items():
        ws.column_dimensions[c].width = w
    put(ws, "B1", "How to use this tracker", F(18, True, INK))
    put(ws, "B2", "A job-search CRM for AI Engineer · Generative AI · Agentic AI · AI Automation · AI Solutions roles.",
        F(9, color=MUTED))
    ws.row_dimensions[1].height = 30
    wrap = Alignment(wrap_text=True, vertical="top", indent=1)
    r = 4

    def block(title, rows, styles=None):
        nonlocal r
        section(ws, f"B{r}", title)
        r += 1
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = put(ws, f"{'BCD'[j]}{r}", val, F(10, j == 0 and len(row) > 1), align=wrap)
                c.border = Border(bottom=side(LINE_SOFT))
            if styles and styles[i]:
                bg, fg = styles[i]
                if bg:
                    ws[f"B{r}"].fill = fill(bg)
                ws[f"B{r}"].font = F(10, True, fg)
            r += 1
        r += 1

    block("Daily routine — 2 minutes", [
        ("1. Open Follow-Ups", "Work top to bottom: overdue first, then due today, no response, due soon, contacts, interviews."),
        ("2. After each action", "On Applications, update Status, Next Action, Next Action Date and Last Updated."),
        ("3. Glance at the Dashboard", "Pipeline, conversion rates and this week's application count."),
        ("Weekly", "Check Analytics: which CV version, source, role and match band actually lead to interviews?"),
    ])
    block("Adding a job", [
        ("Wishlist", "Add a row with Company, Job Title, Job URL, Match %, Status = Wishlist. Copy the 'Next ID' into Application ID."),
        ("When you apply", "Fill Date Applied, set Status = Applied, choose CV Version and Role Category, and set a Next Action + date."),
        ("When they reply", "Enter First Response Date and move Status forward (Screening, Interview 1, …)."),
        ("If it closes", "Set Status = Rejected / Withdrawn. If it got past applying, also set Stage When Closed."),
        ("Interviews", "Add one row per round on Interviews — it feeds Upcoming Interviews on Follow-Ups."),
    ])
    block("Colour legend", [
        ("White cells", "Your input."),
        ("Grey cells", "Auto-calculated — don't type over them."),
        ("Yellow cells (Settings)", "Configurable values: follow-up days, due-soon window, etc."),
    ], styles=[None, (AUTO_BG, MUTED), (INPUT_BG, INPUT_FG)])
    meanings = {
        "Wishlist": "Saved, not applied yet.", "Applied": "Submitted, waiting for a reply.",
        "Screening": "Recruiter / HR screening.", "Interview 1": "First interview round.",
        "Interview 2": "Second interview round.", "Final": "Final round.", "Offer": "Offer received.",
        "Accepted": "Offer accepted.", "Rejected": "Closed by the company — set Stage When Closed if it got past applying.",
        "Withdrawn": "You withdrew — set Stage When Closed if you had applied.",
    }
    block("Status pipeline", [(s, meanings[s]) for s in LISTS["Status"]],
          styles=[(STATUS_STYLE[s][0], STATUS_STYLE[s][1]) for s in LISTS["Status"]])
    block("Action flags (Applications → Action Flag, and the Follow-Ups tab)",
          [(f, m) for f, _, _, m in FLAGS] + [("(none)", "Closed applications (Rejected, Withdrawn, Accepted) never get a flag.")],
          styles=[(bg, fg) for _, bg, fg, _ in FLAGS] + [None])
    block("Match % bands — your own assessment, not an objective score",
          [(f"{cat}", lab) for cat, _, lab, _, _ in MATCH_BANDS], styles=[(b[3], b[4]) for b in MATCH_BANDS])
    block("How the numbers are calculated", [
        ("Total Applications", "Rows whose status is beyond Wishlist (Rejected counts as applied)."),
        ("This Week / This Month", "By Date Applied. Weeks run Monday–Sunday."),
        ("Active", "Applied, Screening, Interview 1/2, Final or Offer — not yet closed."),
        ("Interviews", "Reached Interview 1 or later, including applications later rejected (via Stage When Closed). Screening alone doesn't count."),
        ("Conversion rates", "Application → Interview = Interviews ÷ Total · Interview → Final = Finals ÷ Interviews · "
                             "Final → Offer = Offers ÷ Finals · Overall Offer Rate = Offers ÷ Total."),
        ("Avg. Days to First Response", "Average of First Response Date − Date Applied, for rows that have both."),
        ("Best performer", "Highest interview rate among groups with at least the minimum sample size (Settings)."),
        ("Skill tags", "A tag counts when it appears in Key Technologies / Skills between commas, e.g. 'LLMs, RAG'. Case doesn't matter."),
    ])

    section(ws, f"B{r}", "Example entry — for reference only, not counted anywhere")
    r += 1
    thead(ws, r, [("B", "Column", "left"), ("C", "Example", "left"), ("D", "Tip", "left")])
    r += 1
    for field, example, tip in EXAMPLE:
        put(ws, f"B{r}", field, F(10, True), align=wrap, border=Border(bottom=side(LINE_SOFT)))
        c = put(ws, f"C{r}", example, F(10, color=MUTED if str(example).startswith("(") else TEXT), align=wrap,
                border=Border(bottom=side(LINE_SOFT)))
        if isinstance(example, dt.date):
            c.number_format = DATE_FMT
            c.alignment = Alignment(horizontal="left", vertical="top", indent=1)
        put(ws, f"D{r}", tip, F(9, color=MUTED), align=wrap, border=Border(bottom=side(LINE_SOFT)))
        r += 1
    r += 1
    block("Tips", [
        ("Sorting & filtering", "Use the filter arrows in the header row. Sorting keeps every row intact — IDs travel with their row."),
        ("More than 1,000 jobs?", "Formulas use named ranges (Formulas → Name Manager). Extend the App_* ranges and copy the "
                                  "grey formula columns down."),
        ("Adding dropdown items", "Rename items on Settings freely. To add one, extend that list's range in Data → Data Validation."),
        ("Google Sheets", "File → Import → upload this .xlsx. Dropdowns, colours, formulas and the chart carry over."),
        ("Privacy", "Contacts and notes are personal data — keep the file somewhere private."),
    ])


# ----------------------------------------------------------------------------- sample data (testing only)
def sample_data(today):
    """Dummy records exercising every status, flag and breakdown. Used by verify_tracker.py only."""
    d = lambda n: today + dt.timedelta(days=n)  # noqa: E731
    apps = [
        # id, company, title, status, applied, response, next, closed_at, cv, role, source, match, skills
        (1, "Test Alpha", "AI Engineer", "Applied", d(-10), None, None, None, "AI Engineer CV", "AI Engineer", "LinkedIn", 92, "LLMs, RAG, Python"),
        (2, "Test Beta", "GenAI Engineer", "Applied", d(-3), None, d(-1), None, "Generative AI CV", "Generative AI Engineer", "Company Website", 80, "LLMs; Vector Databases"),
        (3, "Test Gamma", "Agentic AI Engineer", "Screening", d(-20), d(-15), d(0), None, "AI Engineer CV", "Agentic AI Engineer", "Referral", 88, "Agentic AI, AI Agents, Python"),
        (4, "Test Delta", "GenAI Developer", "Interview 1", d(-30), d(-25), d(2), None, "Generative AI CV", "Generative AI Engineer", "LinkedIn", 95, "RAG, LLM APIs"),
        (5, "Test Epsilon", "Automation Engineer", "Interview 2", d(-40), d(-35), d(10), None, "AI Automation CV", "AI Automation Engineer", "Recruiter", 70, "AI Automation, Python"),
        (6, "Test Zeta", "AI Engineer", "Final", d(-50), d(-44), d(1), None, "AI Engineer CV", "AI Engineer", "LinkedIn", 85, "LLMs, RAG"),
        (7, "Test Eta", "AI Solutions Engineer", "Offer", d(-60), d(-55), d(-2), None, "AI Solutions CV", "AI Solutions Engineer", "Referral", 90, "LLMs, SQL"),
        (8, "Test Theta", "AI Engineer", "Accepted", d(-70), d(-60), d(-5), None, "AI Engineer CV", "AI Engineer", "Referral", 93, "Python, RAG"),
        (9, "Test Iota", "Software Engineer", "Rejected", d(-15), d(-10), d(-3), None, "General Software CV", "Other", "Indeed", 55, "SQL, PostgreSQL"),
        (10, "Test Kappa", "GenAI Engineer", "Rejected", d(-25), d(-20), None, "Interview 2", "Generative AI CV", "Generative AI Engineer", "StepStone", 78, "LLMs, RAG, storage"),
        (11, "Test Lambda", "AI Consultant", "Wishlist", None, None, d(5), None, None, "AI Consultant", "Networking", 65, "LLMs"),
        (12, "Test Mu", "Werkstudent KI", "Withdrawn", d(-12), d(-8), None, "Screening", "Werkstudent AI CV", "Werkstudent AI", "University", 60, "Python"),
        (13, "Test Nu", "ML Engineer", "Applied", d(-2), None, None, None, "AI Engineer CV", None, "LinkedIn", None, "Machine Learning"),
        (13, "Test Xi", "AI/ML Engineer", "Applied", d(0), None, None, None, "Custom", "AI/ML Engineer", "LinkedIn", 74, "Machine Learning, Python"),
        (None, "Test Omicron", "AI Engineer", None, None, None, None, None, None, None, None, None, None),
    ]
    interviews = [
        ("Test Delta", "GenAI Developer", "Interview 1", d(2), "Technical", "Not Started"),
        ("Test Epsilon", "Automation Engineer", "Interview 2", d(20), "Hiring Manager", "In Progress"),
        ("Test Zeta", "AI Engineer", "Final", d(-3), "Final", "Completed"),
        ("Test Gamma", "Agentic AI Engineer", "Screening", d(0), "Recruiter", "Ready"),
    ]
    contacts = [
        ("Contact One", "Test Beta", "Recruiter", d(-2), "No Response", "High"),
        ("Contact Two", "Test Gamma", "Engineer", d(2), "Positive", "Medium"),
        ("Contact Three", "Test Delta", "Manager", d(10), "Neutral", "Low"),
        ("Contact Four", "Test Eta", "Alumni", None, None, "Unknown"),
    ]
    companies = [("Test Beta", "AI / ML"), ("Test Theta", "Automotive"), ("Test Nobody", "SaaS")]
    return {"apps": apps, "interviews": interviews, "contacts": contacts, "companies": companies}


def write_sample(wb, today):
    data = sample_data(today)
    ws = wb["Applications"]
    for i, a in enumerate(data["apps"]):
        r = FIRST + i
        (aid, co, title, status, applied, resp, nxt, closed_at, cv, role, source, match, skills) = a
        values = {"A": aid, "B": applied or today, "C": applied, "D": co, "E": title, "J": source, "K": "High",
                  "L": match, "M": status, "N": "Follow up" if nxt else None, "O": nxt, "P": "Recruiter X",
                  "I": "https://example.com/job", "R": cv, "V": skills, "Z": role, "AA": resp, "AB": closed_at}
        for col, v in values.items():
            if v is not None:
                ws[f"{col}{r}"] = v
    ws = wb["Interviews"]
    for i, (co, title, stage, date, typ, prep) in enumerate(data["interviews"]):
        for col, v in zip("ABCDFG", (co, title, stage, date, typ, prep)):
            ws[f"{col}{FIRST + i}"] = v
    ws = wb["Contacts"]
    for i, (name, co, title, fu, resp, pot) in enumerate(data["contacts"]):
        for col, v in zip("ABCJIK", (name, co, title, fu, resp, pot)):
            if v is not None:
                ws[f"{col}{FIRST + i}"] = v
    ws = wb["Companies"]
    for i, (co, ind) in enumerate(data["companies"]):
        ws[f"A{FIRST + i}"], ws[f"B{FIRST + i}"] = co, ind
    return data


# ----------------------------------------------------------------------------- main
def build(path=OUTPUT, sample=False, today=None):
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    names = ["Job Search Dashboard", "Follow-Ups", "Applications", "Interviews", "Companies", "Contacts",
             "Analytics", "Settings", "Guide"]
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(n)
    tabs = {"Job Search Dashboard": "1E293B", "Follow-Ups": "D97706", "Applications": ACCENT,
            "Settings": "94A3B8", "Guide": "94A3B8"}
    for n, color in tabs.items():
        wb[n].sheet_properties.tabColor = color

    refs = {}
    build_settings(wb, wb["Settings"])
    build_applications(wb, wb["Applications"])
    build_interviews(wb, wb["Interviews"])
    build_companies(wb, wb["Companies"])
    build_contacts(wb, wb["Contacts"])
    build_analytics(wb, wb["Analytics"], refs)
    build_followups(wb, wb["Follow-Ups"], refs)
    build_dashboard(wb, wb["Job Search Dashboard"], refs)
    build_guide(wb["Guide"])
    data = write_sample(wb, today or dt.date.today()) if sample else None
    wb.active = 0
    wb.save(path)
    patch_chart_label_format(Path(path))
    return refs, data


def patch_chart_label_format(path):
    """openpyxl writes data-label numFmt without sourceLinked="0", so Excel ignores it and labels empty
    weeks with "0". Add the attribute so the zero-hiding format applies."""
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/charts/chart"):
                data = data.replace(b'<numFmt formatCode="0;-0;;" />', b'<numFmt formatCode="0;-0;;" sourceLinked="0" />')
            dst.writestr(item, data)
    tmp.replace(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sample-data", metavar="PATH", help="write a test copy filled with dummy rows to PATH")
    args = ap.parse_args()
    if args.sample_data:
        build(Path(args.sample_data), sample=True)
        print(f"Test workbook with dummy rows -> {args.sample_data}")
    else:
        build()
        print(f"Clean tracker -> {OUTPUT}")
